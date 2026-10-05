"""Valida e importa imágenes pendientes del catálogo.

Por seguridad, el comportamiento predeterminado solo valida. PostgreSQL se
actualiza únicamente al indicar ``--apply`` y cuando la fila lleva
``aprobado=SI``. El proceso no se ejecuta automáticamente durante un despliegue.
"""

import argparse
import csv
import hashlib
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import httpx
from openpyxl import load_workbook
from sqlalchemy import select

from app.db import SessionLocal
from app.erp_db import fetch_product_image, image_media_type, normalize_image_data
from app.image_missing import mark_product_image_loaded
from app.models import Product, ProductImageMissing


MAX_IMAGE_BYTES = 10 * 1024 * 1024
TRUE_VALUES = {"1", "si", "sí", "s", "true", "yes", "x"}


@dataclass
class ImageInput:
    sku: str
    image_url: str = ""
    image_path: str = ""
    provider: str = ""
    approved: bool = False


def _normalized(row: dict) -> dict[str, str]:
    return {
        str(key or "").strip().lower().replace(" ", "_"): str(value or "").strip()
        for key, value in row.items()
    }


def _input_row(row: dict) -> ImageInput:
    values = _normalized(row)
    sku = (
        values.get("codigo_material")
        or values.get("codigoarticulo")
        or values.get("codigo_articulo")
        or values.get("sku")
        or ""
    )
    return ImageInput(
        sku=sku,
        image_url=values.get("image_url") or values.get("url_imagen") or "",
        image_path=values.get("image_path") or values.get("ruta_imagen") or "",
        provider=values.get("proveedor") or values.get("provider") or "Carga validada",
        approved=(values.get("aprobado") or values.get("approved") or "").lower() in TRUE_VALUES,
    )


def read_input(path: Path) -> list[ImageInput]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.DictReader(handle))
    elif suffix in {".xlsx", ".xlsm"}:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        headers = [str(value or "") for value in next(values, ())]
        rows = [dict(zip(headers, row)) for row in values]
    else:
        raise ValueError("El archivo de imágenes debe ser CSV o XLSX")
    result = []
    for row in rows:
        item = _input_row(row)
        if item.sku:
            result.append(item)
    return result


def _read_image(item: ImageInput, input_dir: Path, client: httpx.Client) -> tuple[bytes, str]:
    if item.image_path:
        image_path = Path(item.image_path).expanduser()
        if not image_path.is_absolute():
            image_path = input_dir / image_path
        raw = image_path.read_bytes()
        source = str(image_path)
    elif item.image_url:
        response = client.get(
            item.image_url,
            headers={"User-Agent": "BermudezUlloa-ImageValidator/1.0"},
        )
        response.raise_for_status()
        raw = response.content
        source = str(response.url)
    else:
        raw = fetch_product_image(item.sku) or b""
        source = "EXITERP SQL Server"
    if not raw:
        raise ValueError("No se encontró una imagen en la fuente indicada")
    if len(raw) > MAX_IMAGE_BYTES:
        raise ValueError("La imagen supera el máximo de 10 MB")
    data = normalize_image_data(raw)
    if not image_media_type(data).startswith("image/"):
        raise ValueError("El archivo no tiene un formato de imagen compatible")
    return data, source


def export_log(output: Path) -> int:
    with SessionLocal() as db:
        rows = db.execute(
            select(ProductImageMissing, Product)
            .join(Product, Product.id == ProductImageMissing.product_id)
            .where(ProductImageMissing.status != "CARGADA")
            .order_by(ProductImageMissing.id)
        ).all()
        with output.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(
                handle,
                fieldnames=(
                    "codigo_material",
                    "nombre_actual",
                    "motivo",
                    "estado",
                    "image_url",
                    "image_path",
                    "proveedor",
                    "aprobado",
                ),
            )
            writer.writeheader()
            for missing, product in rows:
                writer.writerow(
                    {
                        "codigo_material": product.sku,
                        "nombre_actual": product.short_description,
                        "motivo": missing.reason or "",
                        "estado": missing.status,
                        "image_url": "",
                        "image_path": "",
                        "proveedor": "",
                        "aprobado": "NO",
                    }
                )
    print({"status": "ok", "exported": len(rows), "file": str(output)})
    return len(rows)


def process_input(input_path: Path, apply: bool = False) -> dict[str, int]:
    items = read_input(input_path)
    result = {"rows": len(items), "valid": 0, "imported": 0, "not_approved": 0, "failed": 0}
    with SessionLocal() as db, httpx.Client(follow_redirects=True, timeout=30) as client:
        for item in items:
            product = db.scalar(select(Product).where(Product.sku == item.sku))
            if product is None:
                result["failed"] += 1
                print({"sku": item.sku, "status": "error", "detail": "Material inexistente en PostgreSQL"})
                continue
            try:
                data, source = _read_image(item, input_path.parent, client)
                result["valid"] += 1
                if not apply:
                    print({"sku": item.sku, "status": "validada", "source": source, "saved": False})
                    continue
                if not item.approved:
                    result["not_approved"] += 1
                    print({"sku": item.sku, "status": "omitida", "detail": "Falta aprobado=SI"})
                    continue
                product.image_data = data
                product.image_media_type = image_media_type(data)
                product.image_source_url = source if source.startswith(("http://", "https://")) else None
                product.image_source_provider = item.provider or source
                product.image_sha256 = hashlib.sha256(data).hexdigest()
                mark_product_image_loaded(db, product)
                db.commit()
                result["imported"] += 1
                print({"sku": item.sku, "status": "cargada", "source": source})
            except Exception as exc:
                db.rollback()
                result["failed"] += 1
                print({"sku": item.sku, "status": "error", "detail": str(exc)})
    print({"status": "ok", "mode": "apply" if apply else "validation", **result})
    return result


def main(argv: Iterable[str] | None = None) -> None:
    parser = argparse.ArgumentParser(description="Gestiona imágenes pendientes sin cargarlas por defecto")
    parser.add_argument("--input", type=Path, help="CSV/XLSX revisado con las fuentes de imagen")
    parser.add_argument("--export-log", type=Path, help="Exporta el registro de imágenes pendientes a CSV")
    parser.add_argument("--apply", action="store_true", help="Carga solo filas validadas con aprobado=SI")
    args = parser.parse_args(list(argv) if argv is not None else None)
    if args.export_log:
        export_log(args.export_log)
    if args.input:
        process_input(args.input, apply=args.apply)
    if not args.export_log and not args.input:
        parser.error("Indica --export-log o --input")


if __name__ == "__main__":
    main()
