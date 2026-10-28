from datetime import date,timedelta
from math import ceil,sqrt
from statistics import NormalDist
import io
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from app.bootstrap import main as bootstrap
from app.main import app
from app.replenishment_engine import (Parameters,prepare_context,calculate_article,annotate_central,summarize,resolve_parameters)
bootstrap()


def source_for(skus=('03009',),values=None,stock=2):
    values=values if values is not None else [10]*30
    start=date(2026,4,1);end=start+timedelta(days=len(values)-1)
    return {'captured_at':'2026-02-01T10:00:00+00:00','sales_delegations':['00','01','02','04','05'],'company':1,'start':start.isoformat(),'end':end.isoformat(),
        'articles':[{'sku':sku,'description':'Material '+sku,'unit_type':'UD','created_at':'2020-01-01','area_id':1,'family_id':2,'subfamily_id':3,'type_id':4} for sku in skus],
        'sales':[{'sku':sku,'delegation':'02','day':(start+timedelta(days=i)).isoformat(),'sold':value,'returned':0} for sku in skus for i,value in enumerate(values) if value],
        'stocks':[{'sku':sku,'delegation':delegation,'stock':stock if delegation=='02' else 1000,'incoming':0,'outgoing':0} for sku in skus for delegation in ('00','02')]}


def calculate(source,temporary=None):
    context=prepare_context(source,['02'],[],temporary)
    return calculate_article(source['articles'][0],'02',context,include_daily=True)


def test_regular_formula_and_zero_days_and_pending_inventory():
    source=source_for();source['stocks'][1].update(stock=10,incoming=7,outgoing=4)
    row=calculate(source)
    assert row['daily_mean']==10 and row['daily_std']==0 and row['sales_days']==30
    assert row['inventory_position']==13
    assert row['safety_stock']==0 and row['minimum_stock']==10 and row['maximum_stock']==80
    assert row['final_quantity']==67
    assert row['calculation']['z']==pytest.approx(1.6448536269)
    sparse=calculate(source_for(values=[10 if i%2 else 0 for i in range(30)]))
    assert len(sparse['daily_sales'])==30 and sparse['daily_mean']==5
    assert sparse['method']=='PERCENTIL_EMPIRICO_INTERMITENTE'
    assert sparse['maximum_stock']==40


def test_service_level_sigma_and_constraints():
    row=calculate(source_for(values=[9,10,11]*10),{'service_level':99,'max_coverage_days':2,'multiple':7})
    expected=NormalDist().inv_cdf(.99)*row['daily_std']*sqrt(row['parameters']['lead_days'])
    assert row['safety_stock']==pytest.approx(expected) and row['maximum_stock']==20
    assert row['final_quantity']==21
    disabled=calculate(source_for(),{'enabled':False})
    assert disabled['final_quantity']==0
    with pytest.raises(ValueError):Parameters(manual_min=5,manual_max=2)
    with pytest.raises(ValueError):Parameters(multiple=0)
    with pytest.raises(ValueError):Parameters(service_level=100)


def test_no_sales_returns_missing_and_negative_stock():
    source=source_for(values=[0]*30)
    source['sales']=[{'sku':'03009','delegation':'02','day':'2026-04-03','sold':0,'returned':12}]
    row=calculate(source)
    assert row['daily_mean']==0 and row['returns']==12 and row['final_quantity']==0
    assert 'DEVOLUCIONES_ANOMALAS' in [item['code'] for item in row['recommendations']]
    source['stocks']=[]
    row=calculate(source)
    assert row['stock'] is None and row['final_quantity']==0
    assert 'DATOS_STOCK' in [item['code'] for item in row['recommendations']]
    row=calculate(source_for(stock=-2))
    assert row['inventory_position']==-2 and row['final_quantity']==82


def test_comparables_are_low_confidence_and_not_fabricated_history():
    source=source_for(('A','B','C','TARGET'))
    source['sales']=[row for row in source['sales'] if row['sku']!='TARGET']
    context=prepare_context(source,['02'],[])
    row=calculate_article(source['articles'][-1],'02',context)
    assert row['daily_mean']==0 and row['total_demand']==0
    assert row['comparables']['count']==3 and row['minimum_stock']==10 and row['confidence']=='BAJA'
    assert row['method']=='MINIMO_ORIENTATIVO_COMPARABLES'
    assert row['final_quantity']==8
    source['articles']=source['articles'][1:]
    row=calculate_article(source['articles'][-1],'02',prepare_context(source,['02'],[]))
    assert row['comparables'] is None and row['final_quantity']==0


def test_policy_precedence_and_shared_central_shortage():
    article=source_for()['articles'][0]
    policies=[{'scope':scope,'scope_key':key,'parameters':{'review_days':value},'active':True} for scope,key,value in
        [('GLOBAL','*',8),('AREA','1',9),('FAMILY','2',10),('SUBFAMILY','3',11),('TYPE','4',12),('DELEGATION','02',13),('ARTICLE_DELEGATION','03009|02',14)]]
    parameters,rules,origins=resolve_parameters(article,'02',policies)
    assert parameters['review_days']==14 and rules[-1]=='ARTICLE_DELEGATION:03009|02'
    assert origins['review_days']==rules[-1]
    source=source_for();source['stocks'][0]['stock']=100
    row=calculate(source);second={**row,'delegation':'01','recommendations':[]}
    annotate_central([row,second],source)
    assert row['central_total_requested']==156
    assert 'CENTRAL_INSUFICIENTE' in [item['code'] for item in row['recommendations']]
    assert summarize([row,second])['units_proposed']==156


def login(email='operador@bermudez.test'):
    client=TestClient(app)
    assert client.post('/api/v1/auth/login',json={'email':email,'password':'123456'}).status_code==200
    return client


def test_api_calculation_export_and_permissions(monkeypatch):
    import app.replenishment as module
    operator=login();client=login('compras001@cliente.test')
    assert client.get('/api/v1/store/replenishment/options').status_code==403
    raw=source_for()
    monkeypatch.setattr(module,'read_replenishment_source',lambda *args:raw)
    prefix='/api/v1/store/replenishment'
    filters={'date_from':raw['start'],'date_to':raw['end'],'delegations':['02']}
    response=operator.post(prefix+'/simulations',json={'filters':filters})
    assert response.status_code==201,response.text
    first=response.json();sid=first['id']
    assert first['summary']['units_proposed']==78
    rows=operator.get(f'{prefix}/simulations/{sid}/results').json()['items'];rid=rows[0]['id']
    baseline=operator.get(f'{prefix}/simulations/{sid}/results/{rid}').json()
    assert len(baseline['daily_sales'])==30
    recalculated=operator.post(f'{prefix}/simulations/{sid}/results/{rid}/recalculate',json={'parameters':{'review_days':1}}).json()
    assert recalculated['final_quantity']==18
    assert operator.get(f'{prefix}/simulations/{sid}/results/{rid}').json()['final_quantity']==78
    adjusted=operator.post(f'{prefix}/simulations/{sid}/results/{rid}/adjustments',json={'quantity':5,'reason':'Revisión operativa','expected_version':1})
    assert adjusted.status_code==200,adjusted.text
    result=adjusted.json();assert result['result']['calculated_quantity']==78
    assert result['result']['recommended_quantity']==78 and result['result']['adjusted_quantity']==5
    assert result['scenario']['summary']['units_proposed']==5
    assert operator.post(f'{prefix}/simulations/{sid}/results/{rid}/adjustments',json={'quantity':6,'reason':'Edición obsoleta','expected_version':1}).status_code==409
    assert operator.patch(f'{prefix}/simulations/{sid}',json={'name':'Escenario','expected_version':2}).status_code==404
    assert operator.get(prefix+'/compare',params={'a':sid,'b':sid}).status_code==404
    exported=operator.get(f'{prefix}/simulations/{sid}/export.xlsx')
    book=load_workbook(io.BytesIO(exported.content));assert book.active['A2'].value=='03009' and book.active['A2'].data_type=='s'
    assert operator.put(prefix+'/policies',json={'scope':'DELEGATION','scope_key':'02','parameters':{'review_days':5}}).status_code==200
    assert operator.put(prefix+'/policies',json={'scope':'DELEGATION','scope_key':'99','parameters':{}}).status_code==422
    assert operator.get(prefix+'/simulations').status_code==405


def test_source_sql_is_read_only_uses_prompt_fields_and_rejects_duplicate_stock(monkeypatch):
    from contextlib import contextmanager
    import app.replenishment_source as module
    calls=[];duplicate=False
    class Cursor:
        def __enter__(self):return self
        def __exit__(self,*args):pass
        def execute(self,sql,params=()):self.sql=sql;calls.append((sql,params))
        def fetchall(self):
            if 'INFORMATION_SCHEMA' in self.sql:return [{'COLUMN_NAME':name} for name in ('UnidadMedida1_','EX_tipounidad')]
            if 'FROM dbo.articulos' in self.sql:return [{'codigoarticulo':'03009','descripcionarticulo':'Material','replenishment_unit':'ROL','unit_source':'UnidadMedida1_','inactivo':0,'fechaalta':date(2020,1,1),'EX_multiplo':6}]
            if 'AlbaranVentaCabecera' in self.sql:return [{'sku':'03009','delegation':'02','day':date(2026,1,5),'sold':4,'returned':1}]
            rows=[{'sku':'03009','delegation':'00','stock':10,'incoming':2,'outgoing':1}]
            return rows*2 if duplicate else rows
    class Connection:
        def cursor(self):return Cursor()
    @contextmanager
    def connect():yield Connection()
    monkeypatch.setattr(module,'connect_sqlserver',connect)
    result=module.read_replenishment_source(date(2026,1,1),date(2026,1,31),['03009'])
    assert result['articles'][0]['sku']=='03009'
    assert result['articles'][0]['exit_multiple']==6
    assert result['articles'][0]['unit_type']=='ROL' and result['articles'][0]['unit_source']=='UnidadMedida1_'
    assert calls[1][0].index('[UnidadMedida1_]')<calls[1][0].index('[EX_tipounidad]')
    assert 'EX_multiplo' in calls[1][0]
    assert result['sales'][0]['sold']==4 and result['sales'][0]['returned']==1
    assert len(result['stocks'])==1 and result['stocks'][0]['stock']==10
    assert all(sql.strip().upper().startswith('SELECT') for sql,params in calls)
    assert 'avc.CodigoEmpresa=avl.CodigoEmpresa' in calls[2][0]
    assert 'avc.IdDelegacion' in calls[2][0] and 'CodigoAlmacen' not in calls[3][0]
    assert calls[2][1][:2]==(date(2026,1,1),date(2026,2,1))
    duplicate=True
    with pytest.raises(ValueError,match='más de una fila'):module.read_replenishment_source(date(2026,1,1),date(2026,1,31))


def test_source_context_and_history_differences_are_disclosed(monkeypatch):
    import app.replenishment as module
    raw=source_for();monkeypatch.setattr(module,'read_replenishment_source',lambda *args:raw)
    client=login();prefix='/api/v1/store/replenishment';filters={'date_from':raw['start'],'date_to':raw['end'],'delegations':['02']}
    a=client.post(prefix+'/simulations',json={'filters':filters}).json()
    b=client.post(prefix+'/simulations',json={'filters':filters,'source_id':a['id'],'temporary_parameters':{'history_days':10}}).json()
    assert client.get(prefix+'/compare',params={'a':a['id'],'b':b['id']}).status_code==404
    rows=client.get(f"{prefix}/simulations/{a['id']}/results",params={'needed_only':True,'alerts_only':True,'q':'Material','sort':'quantity'}).json()
    assert rows['total']==1 and len(rows['items'])==1

@pytest.mark.parametrize('mode,expected', [('EXIT',84),('MANUAL',80),('NONE',78)])
def test_dispatch_multiple_modes(mode,expected):
    source=source_for();source['articles'][0]['exit_multiple']=12
    row=calculate(source,{'multiple_mode':mode,'multiple':10})
    assert row['quantity_before_rounding']==78
    assert row['final_quantity']==expected
    assert row['effective_multiple']==({'EXIT':12,'MANUAL':10,'NONE':None}[mode])
    # A full dispatch group must not add an extra group.
    source['stocks'][1]['stock']=8
    assert calculate(source,{'multiple_mode':'EXIT'})['final_quantity']==72

@pytest.mark.parametrize('value', [None,0,-1,'bad'])
def test_invalid_exit_multiple_requires_review(value):
    source=source_for();source['articles'][0]['exit_multiple']=value
    row=calculate(source,{'multiple_mode':'EXIT'})
    assert row['final_quantity']==0 and row['effective_multiple'] is None
    assert 'MULTIPLO_EXIT_INVALIDO' in [note['code'] for note in row['recommendations']]
    # The same data must not block an explicitly ungrouped calculation.
    assert calculate(source,{'multiple_mode':'NONE'})['final_quantity']==78


def test_dispatch_days_and_decimal_exit_multiple():
    source=source_for();source['articles'][0]['exit_multiple']=2.5
    row=calculate(source,{'multiple_mode':'EXIT','review_days':3,'lead_days':1})
    assert row['calculation']['horizon']==4 and row['maximum_stock']==40
    assert row['quantity_before_rounding']==38 and row['final_quantity']==40
    assert row['calculation']['dispatch_days']==3 and row['calculation']['lead_days']==1
    row=calculate(source,{'multiple_mode':'EXIT','review_days':1,'lead_days':0,'enabled':False})
    assert row['final_quantity']==0


def test_manual_adjustment_must_respect_exit_dispatch_multiple(monkeypatch):
    import app.replenishment as module
    raw=source_for();raw['articles'][0]['exit_multiple']=12
    monkeypatch.setattr(module,'read_replenishment_source',lambda *args:raw)
    client=login();prefix='/api/v1/store/replenishment'
    scenario=client.post(prefix+'/simulations',json={'filters':{'date_from':raw['start'],'date_to':raw['end'],'delegations':['02']},'temporary_parameters':{'multiple_mode':'EXIT','review_days':7}}).json()
    item=client.get(f"{prefix}/simulations/{scenario['id']}/results").json()['items'][0]
    assert item['final_quantity']==84
    path=f"{prefix}/simulations/{scenario['id']}/results/{item['id']}/adjustments"
    response=client.post(path,json={'quantity':5,'reason':'Prueba de formato','expected_version':1})
    assert response.status_code==422 and 'múltiplo de 12' in response.json()['detail']
    response=client.post(path,json={'quantity':24,'reason':'Dos agrupaciones','expected_version':1})
    assert response.status_code==200 and response.json()['result']['final_quantity']==24


def test_point_keeps_minimum_plus_consumption_until_next_replenishment():
    source=source_for(values=[63.274/30]*30,stock=24)
    row=calculate(source,{'manual_min':21,'review_days':7,'lead_days':1,'multiple_mode':'NONE'})
    assert row['daily_mean']*30==pytest.approx(63.274)
    assert row['expected_cycle_consumption']==pytest.approx(63.274/30*7)
    assert row['replenishment_point']==pytest.approx(21+63.274/30*7) and row['target_stock']==pytest.approx(21+63.274/30*7)
    assert row['final_quantity']==12
    assert row['target_stock']!=row['maximum_stock']
    # The lead time is not added a second time after an explicit minimum.
    assert calculate(source,{'manual_min':21,'review_days':7,'lead_days':5,'multiple_mode':'NONE'})['final_quantity']==12
    source['articles'][0]['exit_multiple']=6
    later=calculate(source,{'manual_min':21,'review_days':14,'multiple_mode':'EXIT'})
    assert later['replenishment_point']==pytest.approx(21+63.274/30*14) and later['quantity_before_rounding']==pytest.approx(21+63.274/30*14-24)
    assert later['final_quantity']==30
    enough=calculate(source_for(values=[63.274/30]*30,stock=40),{'manual_min':21,'review_days':7})
    assert enough['final_quantity']==0


def test_only_explicit_limits_can_reduce_replenishment_point():
    source=source_for(values=[63.274/30]*30,stock=24)
    row=calculate(source,{'manual_min':21,'manual_max':34,'review_days':7,'multiple_mode':'NONE'})
    assert row['replenishment_point']==pytest.approx(21+63.274/30*7) and row['target_stock']==34 and row['final_quantity']==10
    assert 'OBJETIVO_LIMITADO' in [note['code'] for note in row['recommendations']]


def test_article_limit_is_configurable_audited_and_passed_to_exit(monkeypatch):
    import app.replenishment as module
    raw=source_for();calls=[]
    def read(*args):calls.append(args);return dict(raw)
    monkeypatch.setattr(module,'read_replenishment_source',read)
    operator=login();customer=login('compras001@cliente.test');prefix='/api/v1/store/replenishment'
    assert operator.get(prefix+'/settings').json()['article_limit']==20000
    assert customer.put(prefix+'/settings',json={'article_limit':30000}).status_code==403
    assert operator.put(prefix+'/settings',json={'article_limit':100001}).status_code==422
    assert operator.put(prefix+'/settings',json={'article_limit':0}).status_code==422
    response=operator.put(prefix+'/settings',json={'article_limit':30000})
    assert response.status_code==200 and response.json()['article_limit']==30000
    assert operator.get(prefix+'/options').json()['article_limit']==30000
    preview=operator.post(prefix+'/simulations',json={'filters':{'date_from':raw['start'],'date_to':raw['end'],'delegations':['02']},'preview_only':True})
    assert preview.status_code==201 and calls[-1][3]==30000
    operator.put(prefix+'/settings',json={'article_limit':20000})


def test_source_handles_more_than_5000_and_enforces_configured_limit(monkeypatch):
    from contextlib import contextmanager
    import app.replenishment_source as module
    calls=[]
    class Cursor:
        def __enter__(self):return self
        def __exit__(self,*args):pass
        def execute(self,sql,params=()):self.sql=sql;calls.append(sql)
        def fetchall(self):
            if 'FROM dbo.articulos' in self.sql:
                return [{'codigoarticulo':str(i).zfill(5),'descripcionarticulo':'Material','EX_multiplo':6} for i in range(5001)]
            return []
    class Connection:
        def cursor(self):return Cursor()
    @contextmanager
    def connect():yield Connection()
    monkeypatch.setattr(module,'connect_sqlserver',connect)
    result=module.read_replenishment_source(date(2026,1,1),date(2026,1,31),max_articles=6000)
    assert len(result['articles'])==5001 and result['article_limit']==6000
    assert 'TOP 6001' in calls[1]
    with pytest.raises(ValueError,match='límite|más de'):
        module.read_replenishment_source(date(2026,1,1),date(2026,1,31),max_articles=5000)


def test_visible_point_and_maximum_formulas_and_no_consumption_option():
    raw=source_for(values=[63.274/30]*30,stock=24)
    point=calculate(raw,{'simple_formula':True,'manual_min':21,'maximum_days':30,'multiple_mode':'NONE'})
    maximum=calculate(raw,{'simple_formula':True,'manual_min':21,'target_mode':'MAXIMUM','maximum_days':30,'multiple_mode':'NONE'})
    assert point['replenishment_point']==pytest.approx(21+63.274/30*7) and point['maximum_stock']==pytest.approx(84.274)
    assert point['target_stock']==pytest.approx(21+63.274/30*7) and point['final_quantity']==12
    assert maximum['target_stock']==pytest.approx(84.274) and maximum['final_quantity']==61
    assert maximum['calculated_quantity']==pytest.approx(60.274) and maximum['target_mode']=='MAXIMUM'
    assert maximum['calculation_model']=='MAXIMUM_COVERAGE'
    assert 'Stock máximo' in maximum['calculation']['formula']
    excess=calculate(source_for(stock=400),{'simple_formula':True,'maximum_days':30})
    assert excess['overstock'] and excess['excess_stock']==90 and excess['final_quantity']==0
    with pytest.raises(ValueError):Parameters(review_days=14,maximum_days=7)
    # No usar comparables ni una reposición contra stock negativo si se excluyen sin consumo.
    empty=source_for(values=[0]*30,stock=-2)
    excluded=calculate(empty,{'simple_formula':True,'exclude_no_consumption':True,'maximum_days':30})
    assert excluded['no_consumption'] and excluded['final_quantity']==0
    assert 'SIN_CONSUMO_EXCLUIDO' in [r['code'] for r in excluded['recommendations']]
    assert calculate(empty,{'simple_formula':True,'exclude_no_consumption':False,'maximum_days':30})['final_quantity']==2


def test_central_reserves_own_consumption_and_ignores_incoming_for_dispatch():
    raw=source_for();central=raw['stocks'][0];central.update(stock=100,incoming=999,outgoing=10)
    raw['sales'] += [{**row,'delegation':'00','sold':2} for row in list(raw['sales'])]
    row=calculate(raw,{'simple_formula':True,'maximum_days':30,'central_reserve_days':7})
    row['final_quantity']=30
    second={**row,'delegation':'01','recommendations':[]}
    annotate_central([row,second],raw)
    assert row['central_monthly_consumption']==60
    assert row['central_reserve']==16 and row['central_physical_available']==90
    assert row['central_available']==74 and row['central_total_requested']==60
    assert row['central_can_supply'] and second['central_can_supply']
    # La propuesta se comprueba conjuntamente, no se ofrece el saldo completo a cada sede.
    second['final_quantity']=50
    annotate_central([row,second],raw)
    assert row['central_total_requested']==80 and row['central_shortage']==6
    assert not row['central_can_supply'] and not second['central_can_supply']
    assert 'CENTRAL_INSUFICIENTE' in [r['code'] for r in row['recommendations']]
    assert row['final_quantity']==30  # Necesidad no se confunde con despacho autorizado.


def test_central_missing_history_missing_stock_and_negative_data_do_not_allow_dispatch():
    raw=source_for();raw.pop('sales_delegations')
    row=calculate(raw);annotate_central([row],raw)
    assert row['central_available'] is None and not row['central_can_supply']
    assert 'CENTRAL_SIN_DATOS' in [r['code'] for r in row['recommendations']]
    raw['sales_delegations']=['00']
    raw['stocks'][0]['stock']=-1
    annotate_central([row],raw)
    assert row['central_available'] is None and not row['central_can_supply']
    raw['stocks']=raw['stocks'][1:]
    annotate_central([row],raw)
    assert row['central_stock'] is None and not row['central_can_supply']


def test_api_filters_central_overstock_and_zero_consumption(monkeypatch):
    import app.replenishment as module
    raw=source_for(('FILTER-A','FILTER-B','FILTER-C'))
    raw['sales']=[r for r in raw['sales'] if r['sku']!='FILTER-C']
    for stock in raw['stocks']:
        if stock['sku']=='FILTER-B' and stock['delegation']=='00':stock['stock']=1
        if stock['sku']=='FILTER-C' and stock['delegation']=='02':stock['stock']=200
    monkeypatch.setattr(module,'read_replenishment_source',lambda *args:raw)
    client=login();prefix='/api/v1/store/replenishment'
    run=client.post(prefix+'/simulations',json={'filters':{'date_from':raw['start'],'date_to':raw['end'],'delegations':['02']},'temporary_parameters':{'simple_formula':True,'maximum_days':30,'exclude_no_consumption':True}})
    assert run.status_code==201,run.text
    path=prefix+'/simulations/'+run.json()['id']+'/results'
    def skus(**query):return [row['sku'] for row in client.get(path,params=query).json()['items']]
    assert skus(central_only=True)==['FILTER-A']
    assert skus(overstock_only=True)==['FILTER-C']
    assert skus(consumption='WITHOUT')==['FILTER-C']
    assert skus(consumption='WITH')==['FILTER-A','FILTER-B']
    assert skus(needed_only=True)==['FILTER-A','FILTER-B']
    assert skus(needed_only=True,consumption='WITHOUT')==[]
    assert run.json()['summary']['overstock']==1 and run.json()['summary']['central_covered']==1
    exported=client.get(prefix+'/simulations/'+run.json()['id']+'/export.xlsx',params={'mode':'expanded'})
    book=load_workbook(io.BytesIO(exported.content))
    headers=[cell.value for cell in book.active[1]]
    assert 'Consumo mensual Almeiras' in headers and 'Disponible para despacho Almeiras' in headers


def test_filtered_summary_and_export_all_pages_with_exit_control_characters(monkeypatch):
    import app.replenishment as module
    raw=source_for(tuple(f'00{i:03}' for i in range(55)))
    raw['articles'][0]['description']='=Material\x0b de EXIT'
    monkeypatch.setattr(module,'read_replenishment_source',lambda *args:raw)
    client=login();prefix='/api/v1/store/replenishment'
    run=client.post(prefix+'/simulations',json={'filters':{'date_from':raw['start'],'date_to':raw['end'],'delegations':['02']}})
    assert run.status_code==201,run.text
    path=prefix+'/simulations/'+run.json()['id']
    result=client.get(path+'/results',params={'needed_only':True,'page_size':1}).json()
    assert len(result['items'])==1 and result['total']==55
    assert result['summary']['articles_analyzed']==55 and result['summary']['to_replenish']==55
    response=client.get(path+'/export.xlsx',params={'needed_only':True,'sort':'sku','descending':True})
    assert response.status_code==200
    book=load_workbook(io.BytesIO(response.content));sheet=book.active
    assert book.sheetnames==['Reposición']
    headers=[cell.value for cell in sheet[1]]
    assert 'Disponible previsto' in headers and 'Posición' not in headers
    assert 'Seguridad' not in headers and len(headers)==16
    assert sheet.max_row==56 and sheet['A2'].value=='00054'
    assert sheet['A56'].value=='00000' and sheet['A56'].data_type=='s'
    assert sheet['B56'].value=='=Material de EXIT' and sheet['B56'].data_type=='s'
    filtered=client.get(path+'/results',params={'q':'00000'}).json()
    assert filtered['summary']['articles_analyzed']==1
    response=client.get(path+'/export.xlsx',params={'q':'00000'})
    assert load_workbook(io.BytesIO(response.content)).active.max_row==2
    empty=client.get(path+'/results',params={'consumption':'WITHOUT'}).json()
    assert empty['total']==0 and all(value==0 for value in empty['summary'].values())
    response=client.get(path+'/export.xlsx',params={'consumption':'WITHOUT'})
    assert load_workbook(io.BytesIO(response.content)).active.max_row==1


def test_calendar_month_average_estimates_incomplete_month_and_includes_zero_month():
    raw=source_for(values=[0]*30)
    raw.update(start='2026-01-01',end='2026-03-10')
    raw['sales']=[{'sku':'03009','delegation':'02','day':day,'sold':amount,'returned':0} for day,amount in [('2026-01-05',60),('2026-02-05',40),('2026-03-05',20)]]
    row=calculate(raw,{'simple_formula':True,'multiple_mode':'NONE'})
    assert row['monthly_mean']==pytest.approx(54) and row['daily_mean']==pytest.approx(1.8)
    months=row['monthly_consumption']
    assert [m['observed_days'] for m in months]==[31,28,10]
    assert [m['monthly_consumption'] for m in months]==pytest.approx([60,40,62])
    assert [m['estimated'] for m in months]==[False,False,True]
    raw['sales']=[sale for sale in raw['sales'] if not sale['day'].startswith('2026-02')]
    assert calculate(raw)['monthly_mean']==pytest.approx((60+0+62)/3)


def test_partial_start_and_end_leap_month_and_history_window():
    raw=source_for(values=[0]*30)
    raw.update(start='2024-02-15',end='2024-03-10')
    raw['sales']=[{'sku':'03009','delegation':'02','day':day,'sold':amount,'returned':0} for day,amount in [('2024-02-20',15),('2024-03-05',20)]]
    row=calculate(raw)
    assert row['monthly_mean']==pytest.approx((29+62)/2)
    assert row['monthly_consumption'][0]['observed_days']==15
    short=calculate(raw,{'history_days':5})
    assert short['monthly_mean']==0 and short['monthly_consumption'][0]['observed_days']==5


@pytest.mark.parametrize('monthly,stock,outgoing,expected',[(32.787,11,5,3),(1.639,3,0,0)])
def test_rounds_only_final_dispatch_quantity(monthly,stock,outgoing,expected):
    raw=source_for(values=[monthly/30]*30,stock=stock)
    raw['stocks'][1]['outgoing']=outgoing
    row=calculate(raw,{'simple_formula':True,'multiple_mode':'NONE','service_level':50,'lead_days':1,'review_days':7,'maximum_days':30})
    assert row['minimum_stock']==pytest.approx(monthly/30)
    assert row['replenishment_point']==pytest.approx(monthly/30*8)
    assert row['quantity_before_rounding']==pytest.approx(max(0,monthly/30*8-(stock-outgoing)))
    assert row['final_quantity']==expected


def test_statistical_safety_keeps_decimals_until_final_quantity():
    raw=source_for(values=[0,1,2]*10)
    row=calculate(raw,{'simple_formula':True,'multiple_mode':'NONE'})
    expected=NormalDist().inv_cdf(.95)*row['daily_std']*sqrt(row['parameters']['lead_days'])
    assert row['safety_stock']==pytest.approx(expected)
    assert row['minimum_stock']==pytest.approx(row['daily_mean']+expected)
    assert row['final_quantity']==ceil(max(0,row['target_stock']-row['inventory_position']))


def test_review_days_only_change_cycle_consumption_not_statistical_safety():
    source=source_for(values=[0,1,5]*10)
    first=calculate(source,{'simple_formula':True,'review_days':7,'lead_days':1})
    later=calculate(source,{'simple_formula':True,'review_days':14,'lead_days':1})
    assert first['safety_stock']==pytest.approx(later['safety_stock'])
    assert first['minimum_stock']==pytest.approx(later['minimum_stock'])
    assert later['replenishment_point']-first['replenishment_point']==pytest.approx(first['daily_mean']*7)
    slow=calculate(source,{'simple_formula':True,'review_days':7,'lead_days':4})
    assert slow['safety_stock']==pytest.approx(first['safety_stock']*2)
    immediate=calculate(source,{'simple_formula':True,'lead_days':0})
    assert immediate['safety_stock']==0
    source['sales'] += [{**sale,'delegation':'00'} for sale in list(source['sales'])]
    short=calculate(source,{'simple_formula':True,'central_reserve_days':7})
    long=calculate(source,{'simple_formula':True,'central_reserve_days':14})
    annotate_central([short],source);annotate_central([long],source)
    assert long['central_reserve']-short['central_reserve']==pytest.approx(short['central_daily_mean']*7)


def test_monthly_threshold_uses_unrounded_value_and_includes_equal():
    params={'simple_formula':True,'multiple_mode':'NONE','minimum_monthly_consumption':0.5}
    raw=source_for(values=[0.499]+[0]*29,stock=0)
    raw['articles'][0]['unit_type']='ROL'
    below=calculate(raw,params)
    assert below['monthly_mean']==pytest.approx(0.499)
    assert below['final_quantity']==0 and below['below_monthly_threshold']
    assert any(note['code']=='CONSUMO_BAJO_UMBRAL' for note in below['recommendations'])
    assert calculate(raw,{**params,'minimum_monthly_consumption':0})['final_quantity']>0
    raw['sales'][0]['sold']=0.5
    equal=calculate(raw,params)
    assert equal['monthly_mean']==0.5 and not equal['below_monthly_threshold'] and equal['final_quantity']>0
    assert equal['unit_type']=='ROL'
    with pytest.raises(ValueError):Parameters(minimum_monthly_consumption=-1)


@pytest.mark.parametrize('mode',['simple','expanded'])
def test_monthly_threshold_filters_summary_and_excel(monkeypatch,mode):
    import app.replenishment as module
    raw=source_for(('LOW','EQUAL','HIGH'),values=[0]*30,stock=0)
    raw['sales']=[{'sku':sku,'delegation':'02','day':'2026-04-01','sold':quantity,'returned':0} for sku,quantity in [('LOW',0.499),('EQUAL',0.5),('HIGH',0.51)]]
    monkeypatch.setattr(module,'read_replenishment_source',lambda *a:raw)
    client=login();prefix='/api/v1/store/replenishment/simulations'
    response=client.post(prefix,json={'filters':{'date_from':raw['start'],'date_to':raw['end'],'delegations':['02']},'temporary_parameters':{'simple_formula':True,'multiple_mode':'NONE','minimum_monthly_consumption':0.5}})
    assert response.status_code==201,response.text
    path=prefix+'/'+response.json()['id']
    all_rows=client.get(path+'/results').json()
    assert all_rows['total']==3
    assert next(row for row in all_rows['items'] if row['sku']=='LOW')['final_quantity']==0
    filtered=client.get(path+'/results',params={'minimum_monthly_consumption':0.5}).json()
    assert filtered['total']==2 and filtered['summary']['articles_analyzed']==2
    assert {row['sku'] for row in filtered['items']}=={'EQUAL','HIGH'}
    exported=client.get(path+'/export.xlsx',params={'minimum_monthly_consumption':0.5,'mode':mode})
    assert exported.status_code==200
    sheet=load_workbook(io.BytesIO(exported.content)).active
    assert sheet.max_row==3 and {sheet['A2'].value,sheet['A3'].value}=={'EQUAL','HIGH'}
    assert ('Unidad' if mode=='simple' else 'Tipo unidad') in [cell.value for cell in sheet[1]]
    assert client.get(path+'/results',params={'minimum_monthly_consumption':-1}).status_code==422


def test_unit_expression_handles_missing_measurement_columns():
    from app.replenishment_source import unit_expression
    class Cursor:
        def execute(self,*a):pass
        def fetchall(self):return []
    value,origin=unit_expression(Cursor())
    assert value==origin=="'Sin unidad'"
