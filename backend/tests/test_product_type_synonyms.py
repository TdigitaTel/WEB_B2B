import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.bootstrap import main as bootstrap
from app.main import app
from app.db import SessionLocal
from app.models import MaterialArea, MaterialFamily, MaterialSubfamily, MaterialProductType, ProductTypeSynonyms, Product
from app.product_type_synonyms import HEADERS, catalog_search_config

bootstrap()
PREFIX = '/api/v1/store/product-type-synonyms'


def login(email='operador@bermudez.test'):
    client = TestClient(app)
    assert client.post('/api/v1/auth/login', json={'email': email, 'password': '123456'}).status_code == 200
    return client


@pytest.fixture
def types():
    with SessionLocal() as db:
        area=MaterialArea(code='SYN-AREA',name='Sinónimos prueba');db.add(area);db.flush()
        family=MaterialFamily(code='SYN-FAM',name='Familia',area_id=area.id);db.add(family);db.flush()
        sub=MaterialSubfamily(code='SYN-SUB',name='Subfamilia',family_id=family.id);db.add(sub);db.flush()
        first=MaterialProductType(code='000-SYN-01',name='Válvulas de esfera',subfamily_id=sub.id)
        second=MaterialProductType(code='000-SYN-02',name='Válvulas de compuerta',subfamily_id=sub.id)
        db.add_all([first,second]);db.commit();ids=(first.id,second.id,sub.id,family.id,area.id)
    yield ids[:2]
    with SessionLocal() as db:
        for pk in ids[:2]:
            saved=db.get(ProductTypeSynonyms,pk)
            if saved:db.delete(saved)
        db.flush()
        for model,pk in [(MaterialProductType,ids[0]),(MaterialProductType,ids[1]),(MaterialSubfamily,ids[2]),(MaterialFamily,ids[3]),(MaterialArea,ids[4])]:
            db.delete(db.get(model,pk));db.flush()
        db.commit()


def xlsx(rows):
    book=Workbook();sheet=book.active;sheet.title='Sinonimos';sheet.append(HEADERS)
    for code,words in rows:sheet.append([code,'Informativo','','','']+words+['']*(10-len(words)))
    output=io.BytesIO();book.save(output);return output.getvalue()


def upload(client,content,apply=False,mode='merge'):
    return client.post(PREFIX+'/import',files={'file':('sinonimos.xlsx',content)},data={'apply':str(apply).lower(),'mode':mode})


def test_individual_persistence_limits_and_version(types):
    client=login();pk=types[0]
    response=client.put(f'{PREFIX}/{pk}',json={'synonyms':['llave de paso','Llave de corte'],'version':0})
    assert response.status_code==200 and response.json()['version']==1
    item=next(row for row in client.get(PREFIX).json() if row['id']==pk)
    assert item['code']=='000-SYN-01' and item['synonyms']==['llave de paso','Llave de corte']
    assert item['area']=='Sinónimos prueba' and item['family']=='Familia'
    assert client.put(f'{PREFIX}/{pk}',json={'synonyms':['otro'],'version':0}).status_code==409
    for words in [['a']*11,['Válvula','valvula'],[''],['*'],['hola%'],['a'*101]]:
        assert client.put(f'{PREFIX}/{pk}',json={'synonyms':words}).status_code==422
    assert client.put(f'{PREFIX}/{pk}',json={'synonyms':[],'version':1}).json()['synonyms']==[]
    assert client.put(f'{PREFIX}/99999999',json={'synonyms':['test']}).status_code==404


def test_excel_export_roundtrip_and_preview_does_not_write(types):
    client=login();content=xlsx([('000-SYN-01',['llave de paso'])])
    preview=upload(client,content).json()
    assert preview['changed']==1 and preview['applied'] is False and not preview['errors']
    with SessionLocal() as db:assert db.get(ProductTypeSynonyms,types[0]) is None
    assert upload(client,content,True).json()['applied'] is True
    response=client.get(PREFIX+'/export.xlsx')
    assert response.status_code==200 and response.content[:2]==b'PK'
    assert 'sinonimos_tipos_producto.xlsx' in response.headers['content-disposition']
    sheet=load_workbook(io.BytesIO(response.content)).active
    assert sheet.freeze_panes=='A2' and sheet.auto_filter.ref
    row=next(row for row in sheet if row[0].value=='000-SYN-01')
    assert row[0].data_type=='s' and row[5].value=='llave de paso'
    assert upload(client,response.content).json()['changed']==0


def test_bulk_atomic_errors_merge_and_explicit_replace(types):
    client=login();pk=types[0]
    client.put(f'{PREFIX}/{pk}',json={'synonyms':['llave']})
    bad=xlsx([('000-SYN-01',['nuevo']),('NO-EXISTE',['otro'])])
    assert upload(client,bad).json()['errors'][0]['code']=='NO-EXISTE'
    assert upload(client,bad,True).status_code==422
    with SessionLocal() as db:assert db.get(ProductTypeSynonyms,pk).synonyms==['llave']
    assert upload(client,xlsx([('000-SYN-01',['LLAVE','nuevo'])]),True).json()['changed']==1
    with SessionLocal() as db:assert db.get(ProductTypeSynonyms,pk).synonyms==['llave','nuevo']
    assert upload(client,xlsx([('000-SYN-01',[])]),True).json()['changed']==0
    assert upload(client,xlsx([('000-SYN-01',[])]),True,'replace').json()['changed']==1
    with SessionLocal() as db:assert db.get(ProductTypeSynonyms,pk).synonyms==[]
    duplicate=upload(client,xlsx([('000-SYN-01',['a']),('000-SYN-01',['b'])]))
    assert duplicate.json()['errors'][0]['message']=='Tipo repetido en el archivo'


def test_bulk_rejects_more_than_ten_after_merge_and_formulas(types):
    client=login();client.put(f'{PREFIX}/{types[0]}',json={'synonyms':[f'palabra{i}' for i in range(10)]})
    assert upload(client,xlsx([('000-SYN-01',['once'])]),True).status_code==422
    assert upload(client,xlsx([('000-SYN-01',['=SUM(1,2)'])])).status_code==422
    assert client.post(PREFIX+'/import',files={'file':('test.csv',b'csv')}).status_code==422
    assert upload(client,b'invalid zip').status_code==422
    assert upload(client,xlsx([])).status_code==422


def test_synonym_resolution_matches_whole_phrase_and_preserves_extra_terms(types):
    client=login();pk=types[0]
    client.put(f'{PREFIX}/{pk}',json={'synonyms':['llave de paso','llave']})
    with SessionLocal() as db:
        product=db.scalar(select(Product).order_by(Product.id));old=product.material_product_type_id;product.material_product_type_id=pk;db.commit();sku=product.sku;product_id=product.id
        try:
            config=catalog_search_config(db,'LLAVE de PASO latón 1/2')
            assert config['type_synonym']=={'codes':[sku],'query':'laton 1/2'}
            assert 'type_synonym' not in catalog_search_config(db,'llavero')
        finally:
            product=db.get(Product,product_id);product.material_product_type_id=old;db.commit()


@pytest.mark.parametrize('path',[PREFIX,PREFIX+'/export.xlsx'])
def test_permissions(path):
    assert TestClient(app).get(path).status_code==401
    assert login('compras001@cliente.test').get(path).status_code==403


def test_exit_synonyms_preserve_visibility_filters_and_pagination(monkeypatch):
    from app import erp_db
    calls=[]
    class Cursor:
        def __enter__(self):return self
        def __exit__(self,*args):pass
        def execute(self,sql,parameters=()):calls.append((sql,parameters))
        def fetchone(self):return {'total':2}
        def fetchall(self):return []
    class Connection:
        def __enter__(self):return self
        def __exit__(self,*args):pass
        def cursor(self):return Cursor()
    monkeypatch.setattr(erp_db,'connect_sqlserver',Connection)
    monkeypatch.setattr(erp_db,'_brand_columns',lambda cn:{'code':'Marca','name':None,'internet_b2b':'InternetB2B','company':None})
    monkeypatch.setattr(erp_db,'_article_columns',lambda cn:{'code':'CodigoArticulo','description':'DescripcionArticulo','manufacturer_reference':None,'ean':None,'brand_name':None,'brand_code':'Marca','unit':None,'price_with_tax':None,'price_without_tax':None,'inactive':'Inactivo'})
    from app.search import default_search_config
    config=default_search_config();config['type_synonym']={'codes':[f'{i:05}' for i in range(1001)],'query':'laton 1/2'}
    erp_db.fetch_catalog_articles('llave de paso laton 1/2',page=2,article_codes=['00001'],excluded_codes=['00002'],search_config=config)
    inserts=[params for sql,params in calls if sql.startswith('INSERT INTO #SynonymArticles')]
    assert list(map(len,inserts))==[1000,1]
    sql,params=next((sql,params) for sql,params in calls if sql.startswith('SELECT COUNT_BIG'))
    assert '#EligibleArticles' in sql and '#HiddenArticles' in sql and '[Inactivo]' in sql
    assert ' OR (EXISTS (SELECT 1 FROM #SynonymArticles' in sql
    assert params.count('%laton%')==4 and params.count('%1/2%')==4
    assert params.count('%llave%')==2
    assert sql.count(' LIKE %s')==len(params)
    page_sql,page_params=calls[-1]
    assert page_params[:-2]==params and page_params[-2:]==(24,24)


def test_catalogue_and_suggestions_pass_synonym_config(monkeypatch,types):
    client=login();client.put(f'{PREFIX}/{types[0]}',json={'synonyms':['llave de paso']})
    with SessionLocal() as db:
        product=db.scalar(select(Product).order_by(Product.id));pk=product.id;old=product.material_product_type_id;product.material_product_type_id=types[0];db.commit();sku=product.sku
    client=login('compras001@cliente.test')
    calls=[]
    def source(*args,**kwargs):calls.append((args,kwargs));return [],0
    monkeypatch.setattr('app.main.fetch_catalog_articles',source)
    try:
        assert client.get('/api/v1/products',params={'q':'llave de paso'}).status_code==200
        assert client.get('/api/v1/search/suggestions',params={'q':'llave de paso'}).status_code==200
        assert len(calls)==2
        assert all(item[1]['search_config']['type_synonym']=={'codes':[sku],'query':''} for item in calls)
        assert all('excluded_codes' in item[1] for item in calls)
    finally:
        with SessionLocal() as db:db.get(Product,pk).material_product_type_id=old;db.commit()


def test_operator_requires_settings_option(types):
    from app.models import OperatorAccess, User
    client=login()
    with SessionLocal() as db:
        user=db.scalar(select(User).where(User.email=='operador@bermudez.test'))
        access=OperatorAccess(user_id=user.id,exit_code='SYN-PERMISSION-TEST',permissions={'orders':True,'settings':False})
        db.add(access);db.commit();pk=user.id
    try:
        assert client.get(PREFIX).status_code==403
        assert client.get(PREFIX+'/export.xlsx').status_code==403
        assert client.put(f'{PREFIX}/{types[0]}',json={'synonyms':['test']}).status_code==403
        assert upload(client,xlsx([('000-SYN-01',['test'])]),True).status_code==403
    finally:
        with SessionLocal() as db:db.delete(db.get(OperatorAccess,pk));db.commit()
