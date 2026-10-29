"""Write-only Excel tables for web downloads, with literal text cells."""
import io

from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class ExcelTable:
    def __init__(self, title, widths=()):
        self.workbook = Workbook(write_only=True)
        self.sheet = self.workbook.create_sheet(title)
        self.sheet.freeze_panes = "A2"
        self.sheet.row_dimensions[1].height = 32
        for index, width in enumerate(widths, 1):
            self.sheet.column_dimensions[get_column_letter(index)].width = width
        self.rows = 0
        self.columns = 0
        self.header_font = Font(bold=True, color="FFFFFF")
        self.header_fill = PatternFill("solid", fgColor="1C608C")

    def writerow(self, values):
        cells = []
        for value in values:
            if isinstance(value, str):
                cell = WriteOnlyCell(self.sheet, value=ILLEGAL_CHARACTERS_RE.sub("", value))
                # Preserve leading zeros and treat formula-like values as literal text.
                cell.data_type = "s"
                cell.number_format = "@"
            else:
                cell = WriteOnlyCell(self.sheet, value=value)
            if self.rows == 0:
                cell.font = self.header_font
                cell.fill = self.header_fill
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            cells.append(cell)
        self.sheet.append(cells)
        self.rows += 1
        self.columns = max(self.columns, len(cells))

    def response(self, filename):
        if self.rows and self.columns:
            self.sheet.auto_filter.ref = f"A1:{get_column_letter(self.columns)}{self.rows}"
        output = io.BytesIO()
        self.workbook.save(output)
        return Response(
            output.getvalue(),
            media_type=XLSX_MEDIA_TYPE,
            headers={
                "Content-Disposition": f'attachment; filename="{filename}"',
                "Cache-Control": "no-store",
            },
        )
