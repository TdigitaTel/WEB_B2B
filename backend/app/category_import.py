"""Carga manual de la clasificación comercial del catálogo.

SQL Server conserva el maestro de artículos. Este módulo solo administra en
PostgreSQL los metadatos propios de la web: jerarquía, marca, nombre homologado
y criterios de clasificación. Las imágenes quedan fuera de este proceso.
"""

from __future__ import annotations

import csv
import hashlib
import io
import logging
import re
import unicodedata
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy import or_, select
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session

from .auth import audit, require_roles
from .db import get_db
from .erp_db import fetch_catalog_articles_by_codes
from .import_materials import clean, fake_ean, slug
from .models import (
    Brand,
    Category,
    MaterialArea,
    MaterialFamily,
    MaterialImportRow,
    MaterialProductType,
    MaterialSubfamily,
    Product,
    SyncStatus,
    User,
)

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/api/v1/store/catalog-categories", tags=["catalog-categories"])

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ROWS = 5000
SOURCE_SHEET = "CargaCategorias"

TEMPLATE_HEADERS = [
    "CodigoArticulo",
    "Nivel1_Area",
    "Nivel2_Familia",
    "Nivel3_Subfamilia",
    "Nivel4_TipoProducto",
    "Marca",
    "NombreHomologado",
    "CriteriosClasificacion",
    "ReferenciaFabricante",
    "UnidadMedida",
]

FIELD_ALIASES = {
    "code": {"codigoarticulo", "codigomaterial", "codigo", "sku", "referencia"},
    "area": {"nivel1area", "area", "categoria", "departamento"},
    "family": {"nivel2familia", "familia"},
    "subfamily": {"nivel3subfamilia", "subfamilia"},
    "product_type": {"nivel4tipoproducto", "tipoproducto", "tipo"},
    "brand": {"marca", "descripcionmarca"},
    "normalized_name": {
        "nombrehomologado", "nombrenormalizado", "descripcionhomologada",
        "descripcioncomercial", "nombrecomercial",
    },
    "criteria": {
        "criteriosclasificacion", "criterio", "criterios", "observacionesclasificacion",
    },
    "manufacturer_reference": {"referenciafabricante", "referenciafab"},
    "unit": {"unidadmedida", "unidadmedidaventas", "unidad"},
    "status": {"estadoclasificacion", "estado"},
    "confidence": {"confianzaclasificacion", "confianza"},
}

TECHNICAL_ALIASES = {
    "serie_modelo": {"seriemodelo"},
    "sistema_aplicacion": {"sistemaaplicacion"},
    "material": {"material"},
    "medida_diametro": {"medidadiametro", "medida", "diametro"},
    "tipo_conexion": {"tipoconexion", "conexion"},
    "potencia_capacidad": {"potenciacapacidad", "potencia", "capacidad"},
    "formato_presentacion": {"formatopresentacion", "presentacion"},
    "atributos_tecnicos": {"atributostecnicos"},
}


def _header_key(value: Any) -> str:
    text = unicodedata.normalize("NFD", str(value or "").strip().lower())
    plain = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]+", "", plain)


def _text(value: Any, number_format: str | None = None) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, int):
        if number_format and re.fullmatch(r"0+", number_format):
            return str(value).zfill(len(number_format))
        return str(value)
    if isinstance(value, float) and value.is_integer():
        if number_format and re.fullmatch(r"0+", number_format):
            return str(int(value)).zfill(len(number_format))
        return str(int(value))
    return clean(str(value))


def _field_map(headers: list[str]) -> dict[str, str]:
    normalized = {_header_key(header): header for header in headers if clean(header)}
    result: dict[str, str] = {}
    for field, aliases in {**FIELD_ALIASES, **TECHNICAL_ALIASES}.items():
        match = next((normalized[alias] for alias in aliases if alias in normalized), None)
        if match:
            result[field] = match
    return result


def _parse_csv(content: bytes) -> tuple[list[dict[str, str]], str]:
    decoded = None
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            decoded = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if decoded is None:
        raise ValueError("El CSV debe estar codificado en UTF-8 o Windows-1252")
    sample = decoded[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(io.StringIO(decoded), dialect=dialect)
    if not reader.fieldnames:
        raise ValueError("El CSV no contiene una fila de cabeceras")
    rows = [
        {clean(str(key)): _text(value) for key, value in row.items() if key is not None}
        for row in reader
        if any(_text(value) for value in row.values())
    ]
    return rows, "CSV"


def _parse_xlsx(content: bytes) -> tuple[list[dict[str, str]], str]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("No se pudo leer el archivo XLSX") from exc
    try:
        sheet = workbook.active
        iterator = sheet.iter_rows()
        header_cells = next(iterator, None)
        if not header_cells:
            raise ValueError("El XLSX está vacío")
        headers = [_text(cell.value) for cell in header_cells]
        if not any(headers):
            raise ValueError("El XLSX no contiene una fila de cabeceras")
        rows: list[dict[str, str]] = []
        for cells in iterator:
            values = [_text(cell.value, cell.number_format) for cell in cells]
            if not any(values):
                continue
            rows.append({headers[index]: values[index] if index < len(values) else "" for index in range(len(headers))})
        return rows, sheet.title
    finally:
        workbook.close()


def parse_category_file(filename: str, content: bytes) -> tuple[list[dict[str, str]], str]:
    extension = Path(filename).suffix.lower()
    if extension == ".csv":
        return _parse_csv(content)
    if extension in {".xlsx", ".xlsm"}:
        return _parse_xlsx(content)
    raise ValueError("Formato no admitido. Utiliza un archivo CSV o XLSX")


def _canonical_rows(rows: list[dict[str, str]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    if not rows:
        raise ValueError("El archivo no contiene materiales")
    if len(rows) > MAX_ROWS:
        raise ValueError(f"El archivo supera el máximo de {MAX_ROWS} filas")
    mapping = _field_map(list(rows[0]))
    missing_columns = [
        label for field, label in (
            ("code", "CodigoArticulo"),
            ("area", "Nivel1_Area"),
            ("family", "Nivel2_Familia"),
            ("subfamily", "Nivel3_Subfamilia"),
            ("product_type", "Nivel4_TipoProducto"),
        ) if field not in mapping
    ]
    if missing_columns:
        raise ValueError(f"Faltan columnas obligatorias: {', '.join(missing_columns)}")

    valid: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    seen: set[str] = set()
    for index, raw in enumerate(rows, 2):
        item = {field: clean(raw.get(header, "")) for field, header in mapping.items()}
        item["row"] = index
        item["raw"] = raw
        code = item.get("code", "")
        missing = [
            label for field, label in (
                ("code", "código"), ("area", "área"), ("family", "familia"),
                ("subfamily", "subfamilia"), ("product_type", "tipo de producto"),
            ) if not item.get(field)
        ]
        if missing:
            errors.append({"row": index, "code": code or None, "message": f"Falta: {', '.join(missing)}"})
            continue
        if code in seen:
            errors.append({"row": index, "code": code, "message": "Código duplicado en el archivo"})
            continue
        if len(code) > 50:
            errors.append({"row": index, "code": code, "message": "El código supera 50 caracteres"})
            continue
        seen.add(code)
        valid.append(item)
    return valid, errors


def _validate_products(
    db: Session, rows: list[dict[str, Any]], errors: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    """Comprueba en bloques qué códigos necesitan metadatos locales nuevos."""
    existing: set[str] = set()
    codes = [row["code"] for row in rows]
    for start in range(0, len(codes), 500):
        batch = codes[start:start + 500]
        matches = db.execute(
            select(Product.sku, Product.erp_id).where(or_(Product.sku.in_(batch), Product.erp_id.in_(batch)))
        ).all()
        for sku, erp_id in matches:
            if sku:
                existing.add(sku)
            if erp_id:
                existing.add(erp_id)
    valid: list[dict[str, Any]] = []
    for row in rows:
        if row["code"] not in existing and not row.get("normalized_name"):
            errors.append({
                "row": row["row"],
                "code": row["code"],
                "message": "El material no existe en PostgreSQL; indica NombreHomologado para crearlo",
            })
            continue
        valid.append(row)
    return valid


def _validate_exit_articles(
    rows: list[dict[str, Any]], errors: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    """Impide crear metadatos locales para códigos inexistentes o inactivos en EXIT."""
    articles = fetch_catalog_articles_by_codes([row["code"] for row in rows], active_only=True)
    valid: list[dict[str, Any]] = []
    for row in rows:
        article = articles.get(row["code"])
        if not article:
            errors.append({
                "row": row["row"],
                "code": row["code"],
                "message": "El código no existe o está inactivo en EXIT",
            })
            continue
        row["exit_article"] = article
        valid.append(row)
    return valid


def _hierarchy_code(prefix: str, *parts: str) -> str:
    source = "|".join(_header_key(part) for part in parts)
    return f"{prefix}-{hashlib.sha1(source.encode('utf-8')).hexdigest()[:12].upper()}"


def _get_or_create_area(db: Session, name: str) -> tuple[MaterialArea, bool]:
    item = db.scalar(select(MaterialArea).where(MaterialArea.name == name))
    if item:
        return item, False
    item = MaterialArea(code=_hierarchy_code("ARE", name), name=name)
    db.add(item); db.flush()
    return item, True


def _get_or_create_family(db: Session, area: MaterialArea, name: str) -> tuple[MaterialFamily, bool]:
    item = db.scalar(select(MaterialFamily).where(MaterialFamily.area_id == area.id, MaterialFamily.name == name))
    if item:
        return item, False
    item = MaterialFamily(code=_hierarchy_code("FAM", area.name, name), name=name, area_id=area.id)
    db.add(item); db.flush()
    return item, True


def _get_or_create_subfamily(db: Session, family: MaterialFamily, area_name: str, name: str) -> tuple[MaterialSubfamily, bool]:
    item = db.scalar(select(MaterialSubfamily).where(MaterialSubfamily.family_id == family.id, MaterialSubfamily.name == name))
    if item:
        return item, False
    item = MaterialSubfamily(code=_hierarchy_code("SUB", area_name, family.name, name), name=name, family_id=family.id)
    db.add(item); db.flush()
    return item, True


def _get_or_create_type(db: Session, subfamily: MaterialSubfamily, path: tuple[str, str, str], name: str) -> tuple[MaterialProductType, bool]:
    item = db.scalar(select(MaterialProductType).where(MaterialProductType.subfamily_id == subfamily.id, MaterialProductType.name == name))
    if item:
        return item, False
    item = MaterialProductType(code=_hierarchy_code("TIP", *path, name), name=name, subfamily_id=subfamily.id)
    db.add(item); db.flush()
    return item, True


def _get_or_create_brand(db: Session, name: str) -> tuple[Brand, bool]:
    item = db.scalar(select(Brand).where(Brand.name == name))
    if item:
        return item, False
    item = Brand(name=name)
    db.add(item); db.flush()
    return item, True


def _get_or_create_category(db: Session, name: str) -> Category:
    item = db.scalar(select(Category).where(Category.name == name))
    if item:
        return item
    base_slug = f"material-{slug(name)}"[:120] or "material-sin-clasificar"
    candidate = base_slug
    suffix = 1
    while db.scalar(select(Category.id).where(Category.slug == candidate)):
        suffix += 1
        candidate = f"{base_slug[:115]}-{suffix}"
    item = Category(name=name, slug=candidate)
    db.add(item); db.flush()
    return item


def _apply_row(db: Session, row: dict[str, Any], source_file: str, source_sheet: str) -> tuple[str, dict[str, int]]:
    code = row["code"]
    article = row["exit_article"]
    product = db.scalar(select(Product).where(or_(Product.sku == code, Product.erp_id == code)))
    created_product = product is None
    normalized_name = row.get("normalized_name", "")
    if created_product and not normalized_name:
        raise ValueError("El material no existe en PostgreSQL; indica NombreHomologado para crearlo")

    area, area_created = _get_or_create_area(db, row["area"])
    family, family_created = _get_or_create_family(db, area, row["family"])
    subfamily, subfamily_created = _get_or_create_subfamily(db, family, area.name, row["subfamily"])
    product_type, type_created = _get_or_create_type(
        db, subfamily, (area.name, family.name, subfamily.name), row["product_type"],
    )
    brand, brand_created = _get_or_create_brand(db, row.get("brand") or article.get("brand_name") or "SIN MARCA")
    category = _get_or_create_category(db, family.name)

    if product is None:
        product = Product(
            sku=code,
            erp_id=code,
            ean=(article.get("ean") if len(str(article.get("ean") or "")) <= 14 else "") or fake_ean(code),
            manufacturer_reference=row.get("manufacturer_reference") or article.get("manufacturer_reference") or "",
            brand_id=brand.id,
            category_id=category.id,
            family=family.name,
            subfamily=subfamily.name,
            short_description=normalized_name[:240],
            original_description=article.get("description") or None,
            commercial_description=normalized_name,
            technical_description="",
            unit=row.get("unit") or article.get("unit") or "UD",
            pack_size=1,
            list_price=0,
            tax_rate=21,
            attributes={},
            normalized_search="",
            active=True,
            sync_status=SyncStatus.pending,
            source_system="CATEGORY_UPLOAD",
        )
        db.add(product); db.flush()

    name = normalized_name or product.short_description or product.commercial_description or code
    attributes = dict(product.attributes or {})
    if row.get("criteria"):
        attributes["criterios_clasificacion"] = row["criteria"]
    for field in TECHNICAL_ALIASES:
        if row.get(field):
            attributes[field] = row[field]

    product.brand_id = brand.id
    product.category_id = category.id
    product.material_area_id = area.id
    product.material_family_id = family.id
    product.material_subfamily_id = subfamily.id
    product.material_product_type_id = product_type.id
    product.family = family.name
    product.subfamily = subfamily.name
    product.short_description = name[:240]
    product.commercial_description = name
    product.manufacturer_reference = row.get("manufacturer_reference") or product.manufacturer_reference or ""
    product.unit = row.get("unit") or product.unit or "UD"
    product.attributes = attributes
    product.classification_status = row.get("status") or "CLASIFICADO_MANUAL"
    product.classification_reason = row.get("criteria") or product.classification_reason
    product.classification_confidence = row.get("confidence") or product.classification_confidence
    product.source_system = "CATEGORY_UPLOAD"
    product.normalized_search = " ".join(filter(None, [
        code, name, article.get("description"), product.original_description, area.name, family.name,
        subfamily.name, product_type.name, brand.name, *(str(value) for value in attributes.values()),
    ])).lower()
    product.last_sync_at = datetime.now(timezone.utc)

    db.add(MaterialImportRow(
        source_file=source_file,
        source_sheet=source_sheet,
        source_row=row["row"],
        article_code=code,
        product_id=product.id,
        raw_data=row["raw"],
    ))
    return ("created" if created_product else "updated"), {
        "areas": int(area_created),
        "families": int(family_created),
        "subfamilies": int(subfamily_created),
        "product_types": int(type_created),
        "brands": int(brand_created),
    }


@router.get("/template")
def category_template(_: User = Depends(require_roles("OPERADOR_TIENDA", "ADMIN"))):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SOURCE_SHEET
    sheet.append(TEMPLATE_HEADERS)
    sheet.append([
        "01234", "FONTANERÍA", "VÁLVULAS", "VÁLVULAS DE ESFERA", "LATÓN",
        "MARCA", "VÁLVULA ESFERA LATÓN 1/2\"", "Material + medida + tipo de conexión", "REF-123", "UD",
    ])
    sheet.column_dimensions["A"].width = 20
    sheet.column_dimensions["B"].width = 24
    sheet.column_dimensions["C"].width = 28
    sheet.column_dimensions["D"].width = 30
    sheet.column_dimensions["E"].width = 28
    sheet.column_dimensions["F"].width = 20
    sheet.column_dimensions["G"].width = 42
    sheet.column_dimensions["H"].width = 45
    for cell in sheet[1]:
        cell.font = cell.font.copy(bold=True)
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="plantilla_carga_categorias.xlsx"'},
    )


@router.get("/missing.csv")
def missing_category_log(
    _: User = Depends(require_roles("OPERADOR_TIENDA", "ADMIN")),
    db: Session = Depends(get_db),
):
    """Exporta el log de artículos EXIT aún sin metadatos comerciales en PostgreSQL."""
    products = list(db.scalars(
        select(Product).where(
            Product.source_system == "EXIT_CATALOG",
            Product.classification_status == "SIN_CLASIFICAR",
        ).order_by(Product.sku)
    ).all())
    codes = [product.sku for product in products]
    articles: dict[str, dict] = {}
    if codes:
        try:
            articles = fetch_catalog_articles_by_codes(codes, active_only=True)
        except Exception:
            logger.exception("No se pudo enriquecer el log de clasificación con datos de EXIT")

    output = io.StringIO()
    output.write("\ufeff")
    writer = csv.writer(output, delimiter=";")
    writer.writerow(TEMPLATE_HEADERS + ["EstadoVinculo", "Motivo"])
    for product in products:
        article = articles.get(product.sku, {})
        writer.writerow([
            product.sku,
            "", "", "", "",
            article.get("brand_name") or "",
            "",
            "Completar clasificación comercial",
            article.get("manufacturer_reference") or product.manufacturer_reference or "",
            article.get("unit") or product.unit or "UD",
            "SIN_CLASIFICAR",
            "Artículo activo en EXIT sin ficha de categoría vinculada en PostgreSQL",
        ])
    content = output.getvalue().encode("utf-8")
    logger.info("CATALOG_METADATA_MISSING_EXPORT count=%s", len(products))
    return StreamingResponse(
        iter([content]),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": 'attachment; filename="articulos_sin_clasificacion.csv"'},
    )


@router.post("/import")
async def import_categories(
    file: UploadFile = File(...),
    validate_only: bool = Form(False),
    user: User = Depends(require_roles("OPERADOR_TIENDA", "ADMIN")),
    db: Session = Depends(get_db),
):
    filename = Path(file.filename or "carga_categorias").name
    content = await file.read(MAX_FILE_BYTES + 1)
    if not content:
        raise HTTPException(422, "El archivo está vacío")
    if len(content) > MAX_FILE_BYTES:
        raise HTTPException(413, "El archivo supera el máximo de 10 MB")
    try:
        rows, source_sheet = parse_category_file(filename, content)
        valid_rows, errors = _canonical_rows(rows)
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc
    try:
        valid_rows = _validate_exit_articles(valid_rows, errors)
    except Exception as exc:
        logger.exception("No se pudo validar la carga de categorías contra EXIT")
        raise HTTPException(503, "No se pudieron validar los códigos de material en EXIT") from exc
    valid_rows = _validate_products(db, valid_rows, errors)

    result: dict[str, Any] = {
        "filename": filename,
        "validate_only": validate_only,
        "total_rows": len(rows),
        "valid_rows": len(valid_rows),
        "processed": 0,
        "updated": 0,
        "created_products": 0,
        "created": {"areas": 0, "families": 0, "subfamilies": 0, "product_types": 0, "brands": 0},
        "errors": errors,
        "images_imported": 0,
    }
    if validate_only:
        return result

    source_file = f"CATEGORY_UPLOAD/{uuid.uuid4().hex}/{filename[:180]}"
    for row in valid_rows:
        try:
            with db.begin_nested():
                action, created = _apply_row(db, row, source_file, source_sheet)
            result["processed"] += 1
            if action == "created":
                result["created_products"] += 1
            else:
                result["updated"] += 1
            for key, value in created.items():
                result["created"][key] += value
        except (ValueError, SQLAlchemyError) as exc:
            logger.warning("No se pudo clasificar %s desde %s: %s", row.get("code"), filename, exc)
            result["errors"].append({"row": row["row"], "code": row.get("code"), "message": str(exc)})

    audit(
        db,
        user,
        "CATALOG_CATEGORIES_IMPORTED",
        "PRODUCT",
        metadata={
            "filename": filename,
            "processed": result["processed"],
            "errors": len(result["errors"]),
            "created_products": result["created_products"],
            "images_imported": 0,
        },
    )
    db.commit()
    result["valid_rows"] = result["processed"]
    return result
