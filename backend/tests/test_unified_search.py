import io
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from app import main as module, erp_db
from app.db import SessionLocal
from app.models import MaterialArea,MaterialFamily,MaterialSubfamily,MaterialProductType,ProductTypeNameSynonyms,Product,SearchSettings
from app.search import default_search_config
from app.search_resolution import resolve_search, material_groups
from app.product_type_synonyms import HEADERS,catalog_search_config,grouped_types
from app.search_settings import SearchSettingsIn
from app.bootstrap import main as bootstrap
bootstrap()


def groups():
    return [dict(key='accesorio',name='ACCESORIO',is_general=True,synonyms=['accesorios','complemento','complementos'],paths=[]),
            dict(key='machon',name='MACHÓN',is_general=False,synonyms=['niple','niples'],paths=[]),
            dict(key='codo',name='CODO',is_general=False,synonyms=['codos'],paths=[])]


@pytest.mark.parametrize('query,types,terms,ignored',[
    ('accesorio',['ACCESORIO'],[],[]),
    ('complemento',['ACCESORIO'],[],[]),
    ('accesorio de cocina',['ACCESORIO'],[['cocina']],[]),
    ('complemento de cocina',['ACCESORIO'],[['cocina']],[]),
    ('accesorio de latón',[],[['laton']],['accesorio']),
    ('complemento de latón',[],[['laton']],['complemento']),
    ('niple',['MACHÓN'],[],[]),
    ('machón',['MACHÓN'],[],[]),
    ('accesorio niple',['MACHÓN'],[],['accesorio']),
    ('accesorio niple de latón 1/2',['MACHÓN'],[['laton'],['1/2']],['accesorio']),
    ('componentes niple de latón 1/2',['MACHÓN'],[['laton'],['1/2']],['componentes']),
    ('ACCESORIO CODO DE COBRE ½',['CODO'],[['cobre'],['1/2']],['accesorio']),
    ('niiple',[],[['niiple']],[]),
    ('tornillo cocina',[],[['tornillo'],['cocina']],[]),
])
def test_agreed_examples(query,types,terms,ignored):
    plan=resolve_search(query,default_search_config(),groups())
    assert [row['name'] for row in plan['selected_types']]==types
    assert plan['term_groups']==terms and plan['ignored']==ignored and not plan['ambiguous']


def test_longest_alias_material_equivalences_and_ambiguity():
    rows=groups();rows[1]['synonyms']+=['niple de conexión','pieza']
    plan=resolve_search('accesorio niple de conexión INOX 1/2',default_search_config(),rows)
    assert [row['name'] for row in plan['selected_types']]==['MACHÓN']
    assert plan['term_groups']==[['inoxidable','inox'],['1/2']]
    plan=resolve_search('pieza de latón',default_search_config(),rows)
    assert plan['selected_types'][0]['name']=='MACHÓN' and plan['ignored']==[]
    assert resolve_search('niple codo',default_search_config(),rows)['ambiguous']
    rows[2]['synonyms']=['niple']
    assert resolve_search('niple',default_search_config(),rows)['ambiguous']
    config={**default_search_config(),'enabled':False}
    plan=resolve_search('accesorio laton',config,groups())
    assert not plan['selected_types'] and plan['term_groups']==[['laton']]
    assert resolve_search('piezas laton',config,groups())['term_groups']==[['piezas'],['laton']]


def login(email='operador@bermudez.test'):
    client=TestClient(module.app)
    assert client.post('/api/v1/auth/login',json={'email':email,'password':'123456'}).status_code==200
    return client


@pytest.fixture
def classification():
    with SessionLocal() as db:
        keys=['accesorio','machon']
        saved={key:({'synonyms':row.synonyms,'version':row.version,'is_general':row.is_general} if (row:=db.get(ProductTypeNameSynonyms,key)) else None) for key in keys}
        area=MaterialArea(code='UNI-AREA',name='Área unificada');db.add(area);db.flush()
        family=MaterialFamily(code='UNI-FAM',name='Familia unificada',area_id=area.id);db.add(family);db.flush()
        subs=[];types=[]
        for index,name in enumerate(['ACCESORIO','MACHÓN','MACHÓN']):
            sub=MaterialSubfamily(code=f'UNI-SUB-{index}',name=f'Subfamilia {index}',family_id=family.id);db.add(sub);db.flush();subs.append(sub.id)
            item=MaterialProductType(code=f'UNI-TYPE-{index}',name=name,subfamily_id=sub.id);db.add(item);db.flush();types.append(item.id)
        products=list(db.scalars(select(Product).order_by(Product.id).limit(3)));old=[row.material_product_type_id for row in products]
        for row,pk in zip(products,types):row.material_product_type_id=pk
        db.commit();record={'area':area.id,'family':family.id,'subs':subs,'types':types,'products':[row.id for row in products],'skus':[row.sku for row in products]}
    yield record
    with SessionLocal() as db:
        for pk,value in zip(record['products'],old):db.get(Product,pk).material_product_type_id=value
        for key,data in saved.items():
            row=db.get(ProductTypeNameSynonyms,key)
            if data:
                for field,value in data.items():setattr(row,field,value)
            elif row:db.delete(row)
        db.flush()
        for pk in types:db.delete(db.get(MaterialProductType,pk))
        db.flush()
        for pk in subs:db.delete(db.get(MaterialSubfamily,pk))
        db.flush();db.delete(db.get(MaterialFamily,record['family']));db.flush();db.delete(db.get(MaterialArea,record['area']));db.commit()


def configure(client,name,words,general):
    response=client.put('/api/v1/store/product-type-synonyms/group',json={'type_name':name,'synonyms':words,'is_general':general})
    assert response.status_code==200,response.text
    return response.json()


def test_flag_preview_shared_paths_legacy_api_and_excel(classification):
    client=login();path='/api/v1/store/product-type-synonyms'
    configure(client,'ACCESORIO',['complemento'],True)
    saved=configure(client,'MACHÓN',['niple'],False)
    preview=client.get(path+'/preview',params={'q':'complemento niple de latón 1/2'}).json()
    assert preview['selected_types'][0]['name']=='MACHÓN' and preview['ignored']==['complemento']
    with SessionLocal() as db:
        config=catalog_search_config(db,'accesorio niple laton')
        assert set(config['type_synonym']['codes'])==set(classification['skus'][1:])
        assert config['type_synonym']['strict']
    client.put(path+'/group',json={'type_name':'ACCESORIO','synonyms':['complementos']})
    assert next(item for item in client.get(path).json() if item['key']=='accesorio')['is_general'] is True
    configure(client,'MACHÓN',['niple'],True)
    stale=client.put(path+'/group',json={'type_name':'MACHÓN','synonyms':['nuevo'],'is_general':False,'version':saved['version']})
    assert stale.status_code==409
    configure(client,'MACHÓN',['niple'],False)
    response=client.get(path+'/export.xlsx');book=load_workbook(io.BytesIO(response.content));sheet=book['Sinonimos']
    assert sheet.cell(1,len(HEADERS)).value=='General'
    row=next(row for row in sheet if row[0].value=='ACCESORIO');assert row[-1].value=='Sí'
    # A legacy workbook with no General column keeps the current flag.
    old=Workbook();old.active.append(HEADERS[:-1]);old.active.append(['ACCESORIO','']+['complemento']+['']*9)
    output=io.BytesIO();old.save(output)
    uploaded=client.post(path+'/import',files={'file':('old.xlsx',output.getvalue())},data={'apply':'true','mode':'replace'})
    assert uploaded.status_code==200 and uploaded.json()['items'][0]['is_general'] is True
    new=Workbook();new.active.append(HEADERS);new.active.append(['ACCESORIO','']+['complemento']+['']*9+['No'])
    output=io.BytesIO();new.save(output)
    uploaded=client.post(path+'/import',files={'file':('new.xlsx',output.getvalue())},data={'apply':'true','mode':'replace'})
    assert uploaded.status_code==200 and uploaded.json()['items'][0]['is_general'] is False
    assert login('compras001@cliente.test').get(path+'/preview?q=niple').status_code==403


def test_ambiguous_search_uses_422_not_exit_failure(classification):
    operator=login();configure(operator,'ACCESORIO',[],False);configure(operator,'MACHÓN',['niple'],False)
    customer=login('compras001@cliente.test')
    for path in ['/api/v1/products','/api/v1/search/suggestions']:
        response=customer.get(path,params={'q':'accesorio niple'})
        assert response.status_code==422,response.text
        assert 'varios tipos' in response.json()['detail']


def test_material_groups_validation_and_legacy_settings():
    from app.search_settings import get_search_config
    before=login().get('/api/v1/store/search-settings').json();client=login()
    body={**default_search_config(),'material_groups':[{'name':'Inoxidable','synonyms':['inox','acero inoxidable']},{'name':'Latón','synonyms':[]}]}
    try:
        response=client.patch('/api/v1/store/search-settings',json=body)
        assert response.status_code==200,response.text
        with SessionLocal() as db:config=get_search_config(db)
        plan=resolve_search('accesorio acero inoxidable 1/2',config,groups())
        assert not plan['selected_types'] and plan['term_groups']==[['inoxidable','inox','acero inoxidable'],['1/2']]
        for aliases in [[{'name':'A','synonyms':['x']},{'name':'B','synonyms':['x']}],[{'name':'A','synonyms':['bad%']}],[{'name':'A','synonyms':['bad_']}]]:
            assert client.patch('/api/v1/store/search-settings',json={**body,'material_groups':aliases}).status_code==422
        legacy={**default_search_config(),'materials':['cobre']}
        assert client.patch('/api/v1/store/search-settings',json=legacy).status_code==200
        assert client.get('/api/v1/store/search-settings').json()['material_groups']==[{'name':'cobre','synonyms':[]}]
    finally:client.patch('/api/v1/store/search-settings',json=before)


def test_strict_sql_and_material_or_preserve_visibility(monkeypatch):
    calls=[]
    class Cursor:
        def __enter__(self):return self
        def __exit__(self,*args):pass
        def execute(self,sql,parameters=()):calls.append((sql,parameters))
        def fetchone(self):return {'total':2}
        def fetchall(self):return []
    class Connection:
        def __enter__(self):return self
        def __exit__(self,*args):pass
        def cursor(self):return Cursor()
    monkeypatch.setattr(erp_db,'connect_sqlserver',Connection)
    monkeypatch.setattr(erp_db,'_brand_columns',lambda cn:{'code':'Marca','name':None,'internet_b2b':'InternetB2B','company':None})
    monkeypatch.setattr(erp_db,'_article_columns',lambda cn:{'code':'CodigoArticulo','description':'DescripcionArticulo','manufacturer_reference':None,'ean':None,'brand_name':None,'brand_code':'Marca','unit':None,'price_with_tax':None,'price_without_tax':None,'inactive':'Inactivo'})
    plan=resolve_search('accesorio niple inox 1/2',default_search_config(),groups())
    config={**default_search_config(),'search_plan':plan,'type_synonym':{'codes':['MACHON-1'],'strict':True,'query':plan['query']}}
    erp_db.fetch_catalog_articles('accesorio niple inox 1/2',page=2,article_codes=['MACHON-1'],excluded_codes=['HIDDEN'],search_config=config)
    sql,params=next((sql,params) for sql,params in calls if sql.startswith('SELECT COUNT_BIG'))
    assert '#EligibleArticles' in sql and '#HiddenArticles' in sql and '[Inactivo]' in sql and '#SynonymArticles' in sql
    assert ' OR (EXISTS' not in sql
    assert '%accesorio%' not in params and '%niple%' not in params
    assert '%inox%' in params and '%inoxidable%' in params and '%1/2%' in params
    assert sql.count(' LIKE %s')==len(params)
    assert calls[-1][1][-2:]==(24,24)
    calls.clear();config['type_synonym']['codes']=[]
    erp_db.fetch_catalog_articles('niple',search_config=config)
    sql=next(sql for sql,params in calls if sql.startswith('SELECT COUNT_BIG'))
    assert '1=0' in sql


def test_empty_catalogue_query_does_not_load_type_hierarchy(monkeypatch):
    from app import product_type_synonyms as module
    monkeypatch.setattr(module,'get_search_config',lambda db:default_search_config())
    def forbidden(*args):raise AssertionError('No hace falta cargar tipos para listar el catálogo sin búsqueda')
    monkeypatch.setattr(module,'grouped_types',forbidden)
    assert module.catalog_search_config(None,'')['search_plan']['term_groups']==[]
