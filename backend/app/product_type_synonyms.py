"""One synonym list per normalized type name, shared by all classification paths."""
import hashlib
import io
import json
import re
import unicodedata
from zipfile import ZipFile

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from openpyxl import load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from pydantic import BaseModel, Field, field_validator
from sqlalchemy import select, text
from sqlalchemy.orm import Session

from .auth import audit, require_roles
from .db import get_db
from .excel_download import ExcelTable
from .models import (MaterialArea, MaterialFamily, MaterialSubfamily, MaterialProductType,
                     ProductTypeSynonyms, ProductTypeNameSynonyms, Product, User)
from .search import normalize_query
from .search_settings import get_search_config

router = APIRouter(prefix="/api/v1/store/product-type-synonyms", tags=["product-type-synonyms"])
HEADERS = ["TipoProducto", "CadenasDeClasificacion"] + [f"Sinonimo{i}" for i in range(1, 11)]
MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 10000
operator = require_roles("OPERADOR_TIENDA", "ADMIN")

def alias_key(value):
    return " ".join("".join(c for c in unicodedata.normalize("NFKD", value.casefold()) if not unicodedata.combining(c)).split())


class SynonymsIn(BaseModel):
    synonyms: list[str] = Field(max_length=10)
    version: str | None = Field(default=None, max_length=64)

    @field_validator("synonyms")
    @classmethod
    def validate_synonyms(cls, values):
        result, seen = [], set()
        for value in values:
            value = " ".join(value.split())
            if not value or len(value) > 100 or not any(c.isalnum() for c in value):
                raise ValueError("Cada sinónimo debe tener entre 1 y 100 caracteres y contener letras o números")
            if any(unicodedata.category(c).startswith("C") for c in value) or any(c in value for c in "%_*"):
                raise ValueError("Los sinónimos no admiten controles ni comodines %, _ o *")
            key = alias_key(value)
            if key in seen:
                raise ValueError("Hay sinónimos repetidos (se ignoran mayúsculas y acentos)")
            seen.add(key)
            result.append(value)
        return result


def type_rows(db):
    return db.execute(select(MaterialProductType, MaterialSubfamily.name, MaterialFamily.name, MaterialArea.name, ProductTypeSynonyms)
        .join(MaterialSubfamily, MaterialProductType.subfamily_id == MaterialSubfamily.id)
        .join(MaterialFamily, MaterialSubfamily.family_id == MaterialFamily.id)
        .join(MaterialArea, MaterialFamily.area_id == MaterialArea.id)
        .outerjoin(ProductTypeSynonyms, ProductTypeSynonyms.product_type_id == MaterialProductType.id)
        .order_by(MaterialArea.name, MaterialFamily.name, MaterialSubfamily.name, MaterialProductType.name)).all()


class GroupSynonymsIn(SynonymsIn):
    type_name: str = Field(min_length=1, max_length=160)


def grouped_types(db):
    common = {row.type_name_key: row for row in db.scalars(select(ProductTypeNameSynonyms).execution_options(populate_existing=True)).all()}
    groups = {}
    for product_type, subfamily, family, area, old in type_rows(db):
        key = alias_key(product_type.name)
        group = groups.setdefault(key, {"key": key, "name": product_type.name, "paths": [], "synonyms": [], "legacy_versions": []})
        group["paths"].append({"id": product_type.id, "code": product_type.code, "area": area,
                               "category_family": family, "subfamily": subfamily})
        group["legacy_versions"].append((product_type.id, old.version if old else 0))
        for alias in old.synonyms if old else []:
            if alias_key(alias) not in {alias_key(item) for item in group["synonyms"]}:
                group["synonyms"].append(alias)
    for group in groups.values():
        saved = common.get(group["key"])
        if saved is not None:
            group["synonyms"] = saved.synonyms
        # Preserve all legacy words, even if the combined history exceeds ten.
        group["needs_review"] = len(group["synonyms"]) > 10
        state = [group["key"], saved.version if saved else 0, sorted(group["legacy_versions"]) if saved is None else [], sorted(path["id"] for path in group["paths"])]
        group["version"] = hashlib.sha256(json.dumps(state, ensure_ascii=False).encode()).hexdigest()
        del group["legacy_versions"]
    return sorted(groups.values(), key=lambda item: item["key"])


def lock_groups(db, keys):
    if db.bind.dialect.name == "postgresql":
        for key in sorted(set(keys)):
            db.execute(text("SELECT pg_advisory_xact_lock(hashtextextended(:key, 0))"), {"key": "type-name-synonyms:" + key})


def save_group(db, group, aliases):
    row = db.get(ProductTypeNameSynonyms, group["key"], populate_existing=True)
    if row is None:
        row = ProductTypeNameSynonyms(type_name_key=group["key"], version=0)
        db.add(row)
    row.synonyms = aliases
    row.version += 1
    return row


@router.get("")
def list_types(user: User = Depends(operator), db: Session = Depends(get_db)):
    return grouped_types(db)


def chain_description(path):
    return f"Área: {path['area']} → Categoría / familia: {path['category_family']} → Subfamilia: {path['subfamily']}"


@router.get("/export.xlsx")
def export_types(user: User = Depends(operator), db: Session = Depends(get_db)):
    groups = grouped_types(db)
    if any(group["needs_review"] for group in groups):
        raise HTTPException(409, "Hay tipos con más de 10 sinónimos previos. Revisa sus listas antes de descargar el Excel")
    table = ExcelTable("Sinonimos", [32, 100] + [24] * 10)
    table.writerow(HEADERS)
    for group in groups:
        table.writerow([group["name"], "\n".join(chain_description(path) for path in group["paths"])]
                       + group["synonyms"] + [""] * (10 - len(group["synonyms"])))
    sheet = table.workbook.create_sheet("Cadenas")
    sheet.freeze_panes = "A2"
    sheet.append(["TipoProducto", "Area", "Categoria_Familia", "Subfamilia"])
    for group in groups:
        for path in group["paths"]:
            # Literal text, including names beginning with a formula marker.
            cells = []
            for value in [group["name"], path["area"], path["category_family"], path["subfamily"]]:
                cell = WriteOnlyCell(sheet, value=ILLEGAL_CHARACTERS_RE.sub("", value)); cell.data_type = "s"; cells.append(cell)
            sheet.append(cells)
    return table.response("sinonimos_tipos_producto.xlsx")


@router.put("/group")
def update_group(data: GroupSynonymsIn, user: User = Depends(operator), db: Session = Depends(get_db)):
    key = alias_key(data.type_name)
    lock_groups(db, [key])
    group = next((item for item in grouped_types(db) if item["key"] == key), None)
    if group is None:
        raise HTTPException(404, "Tipo de producto no encontrado")
    if data.version is not None and data.version != group["version"]:
        raise HTTPException(409, "Los sinónimos o las cadenas cambiaron. Actualiza la lista antes de guardar")
    save_group(db, group, data.synonyms)
    audit(db, user, "PRODUCT_TYPE_SYNONYMS_UPDATED", "MATERIAL_PRODUCT_TYPE_NAME", hashlib.sha256(key.encode()).hexdigest(),
          {"type_name": group["name"], "synonyms": data.synonyms, "paths": len(group["paths"])})
    db.commit()
    return next(item for item in grouped_types(db) if item["key"] == key)


@router.put("/{product_type_id}")
def update_type(product_type_id: int, data: SynonymsIn, user: User = Depends(operator), db: Session = Depends(get_db)):
    """Old integrations can identify a path, but now update the shared name group."""
    parent = db.get(MaterialProductType, product_type_id)
    if parent is None:
        raise HTTPException(404, "Tipo de producto no encontrado")
    return update_group(GroupSynonymsIn(type_name=parent.name, synonyms=data.synonyms, version=data.version), user, db)


def read_import(content):
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError("El contenido descomprimido del Excel supera 50 MB")
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        try:
            sheet = book["Sinonimos"] if "Sinonimos" in book.sheetnames else book.active
            if sheet.max_column and sheet.max_column > 30:
                raise ValueError("El Excel contiene demasiadas columnas. Usa la plantilla descargada")
            rows = sheet.iter_rows()
            first = next(rows, ())
            headers = [str(cell.value or "").strip() for cell in first]
            required = [f"Sinonimo{i}" for i in range(1, 11)]
            if any(headers.count(key) != 1 for key in required) or not (headers.count("TipoProducto") == 1 or headers.count("CodigoTipoProducto") == 1):
                raise ValueError("Usa el Excel descargado: TipoProducto y columnas Sinonimo1 a Sinonimo10")
            result = []
            for number, cells in enumerate(rows, 2):
                if number > MAX_ROWS + 1:
                    raise ValueError(f"El archivo admite hasta {MAX_ROWS} filas de tipos")
                values = {header: cells[i].value if i < len(cells) else None for i, header in enumerate(headers)}
                if all(value is None or str(value).strip() == "" for value in values.values()):
                    continue
                if any(cell.data_type == "f" for cell in cells):
                    raise ValueError(f"Fila {number}: escribe valores, no fórmulas")
                result.append((number, values))
            return result
        finally:
            book.close()
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, str(exc) if isinstance(exc, ValueError) else "No se pudo leer el Excel. Usa un archivo .xlsx válido") from exc


@router.post("/import")
def import_types(file: UploadFile = File(...), mode: str = Form("merge"), apply: bool = Form(False),
                 user: User = Depends(operator), db: Session = Depends(get_db)):
    if mode not in {"merge", "replace"}:
        raise HTTPException(422, "Elige añadir o reemplazar sinónimos")
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(422, "Carga un archivo Excel .xlsx")
    content = file.file.read(MAX_FILE_BYTES + 1)
    if len(content) > MAX_FILE_BYTES:
        raise HTTPException(413, "El archivo supera 5 MB")
    imported = read_import(content)
    if not imported:
        raise HTTPException(422, "El Excel no contiene tipos de producto")
    groups = {item["key"]: item for item in grouped_types(db)}
    by_code = {path["code"]: group["key"] for group in groups.values() for path in group["paths"]}
    def row_key(values):
        code = str(values.get("CodigoTipoProducto") or "").strip()
        return by_code.get(code, "") if code else alias_key(str(values.get("TipoProducto") or ""))
    if apply:
        lock_groups(db, [row_key(values) for _, values in imported if row_key(values) in groups])
        groups = {item["key"]: item for item in grouped_types(db)}
    errors, seen, plan = [], set(), []
    for number, values in imported:
        key = row_key(values)
        name = str(values.get("TipoProducto") or values.get("CodigoTipoProducto") or "").strip()
        if key in seen:
            errors.append({"row": number, "name": name, "message": "Tipo repetido. Usa una sola fila por nombre de tipo, aunque tenga varias cadenas"}); continue
        seen.add(key)
        group = groups.get(key)
        if group is None:
            errors.append({"row": number, "name": name, "message": "Tipo de producto desconocido. Usa el nombre de la plantilla"}); continue
        incoming = [str(values[f"Sinonimo{i}"]).strip() for i in range(1, 11) if values.get(f"Sinonimo{i}") is not None and str(values[f"Sinonimo{i}"]).strip()]
        try:
            incoming = SynonymsIn(synonyms=incoming).synonyms
            combined = list(group["synonyms"]) if mode == "merge" else []
            for value in incoming:
                if alias_key(value) not in {alias_key(item) for item in combined}:
                    combined.append(value)
            combined = SynonymsIn(synonyms=combined).synonyms
            plan.append({"key": key, "name": group["name"], "synonyms": combined, "paths": group["paths"],
                         "changed": combined != group["synonyms"]})
        except ValueError:
            errors.append({"row": number, "name": name, "message": "Máximo 10 sinónimos compartidos, sin repetidos ni comodines; cada uno de 1 a 100 caracteres"})
    if apply and errors:
        raise HTTPException(422, {"message": "No se guardó ningún cambio. Corrige el Excel", "errors": errors})
    if apply:
        for item in plan:
            if item["changed"]:
                save_group(db, groups[item["key"]], item["synonyms"])
        audit(db, user, "PRODUCT_TYPE_SYNONYMS_IMPORTED", "MATERIAL_PRODUCT_TYPE_NAME",
              metadata={"mode": mode, "types": len(plan), "changed": sum(item["changed"] for item in plan)})
        db.commit()
    return {"applied": apply, "total": len(imported), "changed": sum(item["changed"] for item in plan), "errors": errors, "items": plan}


def catalog_search_config(db, query):
    config = get_search_config(db)
    normalized = normalize_query(query)
    if not normalized:
        return config
    matches = []
    groups = grouped_types(db)
    for group in groups:
        for alias in group["synonyms"]:
            phrase = normalize_query(alias)
            match = re.search(r"(?<!\w)" + re.escape(phrase) + r"(?!\w)", normalized) if phrase else None
            if match:
                matches.append((len(phrase), phrase, match.start(), match.end(), group["key"]))
    if not matches:
        return config
    best = max(matches, key=lambda item: (item[0], -item[2]))
    keys = {item[4] for item in matches if item[1:4] == best[1:4]}
    ids = {path["id"] for group in groups if group["key"] in keys for path in group["paths"]}
    codes = list(db.scalars(select(Product.sku).where(Product.material_product_type_id.in_(ids))).all())
    if codes:
        config["type_synonym"] = {"codes": codes, "query": (normalized[:best[2]] + " " + normalized[best[3]:]).strip()}
    return config
