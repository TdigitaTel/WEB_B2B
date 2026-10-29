"""API de simulación de reposición. Nunca genera outbox ni escribe en EXIT."""
import hashlib
import io
import json
import logging
import math
from datetime import date, datetime
from typing import Literal
from zoneinfo import ZoneInfo
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field, model_validator
from sqlalchemy import select, func, or_, case
from sqlalchemy.orm import Session
from .auth import require_roles
from .db import get_db
from .models import (AuditEvent, MaterialArea, MaterialFamily, MaterialSubfamily, MaterialProductType, Product,
    ReplenishmentPolicy as Policy, ReplenishmentScenario as Scenario, ReplenishmentResult as Result,
    ReplenishmentAdjustment as Adjustment, ReplenishmentSettings as ModuleSettings,
    ReplenishmentJob as Job, ReplenishmentBatch as Batch, User)
from .replenishment_engine import Parameters, PRECEDENCE, prepare_context, calculate_article, annotate_central, summarize, rounded_dispatch
from .replenishment_source import DELEGATIONS, MAX_ARTICLES, HARD_MAX_ARTICLES, read_replenishment_source

operator=require_roles('ADMIN','OPERADOR_TIENDA')
router=APIRouter(prefix='/api/v1/store/replenishment',tags=['replenishment'],dependencies=[Depends(operator)])
logger=logging.getLogger(__name__)

def today():return datetime.now(ZoneInfo('Europe/Madrid')).date()

def validate_parameters(patch):
    try:Parameters.model_validate({**Parameters().model_dump(),**patch})
    except ValueError as exc:raise HTTPException(422,str(exc)) from exc
    return patch

class Filters(BaseModel):
    date_from: date = date(2026,1,1)
    date_to: date = Field(default_factory=today)
    delegations: list[Literal['01','02','04','05']] = Field(default_factory=lambda:list(DELEGATIONS),min_length=1,max_length=4)
    area_id:int|None=None
    family_id:int|None=None
    subfamily_id:int|None=None
    type_id:int|None=None
    skus:list[str]=Field(default_factory=list,max_length=HARD_MAX_ARTICLES)
    @model_validator(mode='after')
    def dates(self):
        if self.date_from>self.date_to or self.date_to>today() or (self.date_to-self.date_from).days>1095:
            raise ValueError('El periodo debe terminar como máximo hoy y abarcar entre 1 y 1.096 días')
        self.delegations=list(dict.fromkeys(self.delegations))
        self.skus=list(dict.fromkeys(code.strip() for code in self.skus if code.strip()))
        if any(len(code)>50 for code in self.skus):raise ValueError('Código de artículo demasiado largo')
        return self

class PolicyInput(BaseModel):
    scope:Literal['GLOBAL','AREA','FAMILY','SUBFAMILY','TYPE','DELEGATION','ARTICLE_DELEGATION']
    scope_key:str=Field(min_length=1,max_length=100)
    parameters:dict=Field(default_factory=dict)
    active:bool=True

class RunInput(BaseModel):
    filters:Filters=Field(default_factory=Filters)
    temporary_parameters:dict=Field(default_factory=dict)
    source_id:str|None=None
    preview_only:bool=False

class RecalculateInput(BaseModel):
    parameters:dict=Field(default_factory=dict)

class AdjustmentInput(BaseModel):
    parameters:dict=Field(default_factory=dict)
    quantity:float|None=Field(None,ge=0,le=1e9,allow_inf_nan=False)
    excluded:bool=False
    reason:str=Field(min_length=3,max_length=500)
    expected_version:int=Field(ge=1)

def policy_payload(row):
    return {'id':row.id,'scope':row.scope,'scope_key':row.scope_key,'parameters':row.parameters,'active':row.active}

def scenario_or_404(db,public_id,lock=False):
    query=select(Scenario).where(Scenario.public_id==public_id)
    if lock:query=query.with_for_update()
    row=db.scalar(query)
    if not row:raise HTTPException(404,'Simulación no encontrada')
    ensure_published(db,row.id)
    return row

def ensure_published(db,scenario_id):
    status=db.scalar(select(Job.status).where(Job.scenario_id==scenario_id))
    if status and status!='COMPLETED':raise HTTPException(409,'El cálculo todavía no está completo. Consulta su progreso.')

def result_or_404(db,scenario,result_id):
    row=db.scalar(select(Result).where(Result.id==result_id,Result.scenario_id==scenario.id))
    if not row:raise HTTPException(404,'Artículo de simulación no encontrado')
    return row

def audit(db,user,action,entity,metadata):
    db.add(AuditEvent(user_id=user.id,action=action,entity_type='REPLENISHMENT',entity_id=str(entity),metadata_json=metadata))

def scenario_payload(row):
    return {'id':row.public_id,'name':row.name,'saved':row.saved,'version':row.version,'created_at':row.created_at,
        'captured_at':row.source_snapshot['captured_at'],'extraction_started_at':row.source_snapshot.get('extraction_started_at'),'source_fingerprint':row.source_fingerprint,
        'filters':row.filters,'temporary_parameters':row.temporary_parameters,'summary':row.summary,
        'article_limit':row.source_snapshot.get('article_limit',5000),'source_articles':row.source_snapshot.get('source_articles',len(row.source_snapshot['articles'])),'read_only_exit':True,
        'source_notes':['Stock transferible Almeiras = stock físico − comprometido − reserva de consumo; no incluye entradas pendientes.',
            'Ventas: empresa 1. La vista de stock se usa con el alcance indicado, sin campo de empresa.',
            'Ventas del periodo elegido y stock consultado en la extracción; no se reconstruye stock histórico.',
            'El vínculo cabecera/líneas usa número, ejercicio y empresa según el esquema facilitado.']}

class SettingsInput(BaseModel):
    article_limit:int=Field(ge=1,le=HARD_MAX_ARTICLES)

def module_settings(db):
    row=db.get(ModuleSettings,1)
    return {'article_limit':row.article_limit if row else MAX_ARTICLES,'maximum_article_limit':HARD_MAX_ARTICLES}

@router.get('/settings')
def read_settings(db:Session=Depends(get_db)):
    return module_settings(db)

@router.put('/settings')
def save_settings(body:SettingsInput,db:Session=Depends(get_db),user:User=Depends(operator)):
    previous=module_settings(db)
    row=db.get(ModuleSettings,1)
    if row is None:
        row=ModuleSettings(id=1);db.add(row)
    row.article_limit=body.article_limit;row.updated_by=user.id
    audit(db,user,'REPLENISHMENT_SETTINGS_SAVED','1',{'previous':previous,'article_limit':body.article_limit})
    db.commit();return module_settings(db)

@router.get('/options')
def options(db:Session=Depends(get_db)):
    result={'server_date':today().isoformat(),'delegations':[{'code':key,'name':value} for key,value in DELEGATIONS.items()],
        'defaults':Parameters().model_dump(),'precedence':PRECEDENCE,**module_settings(db)}
    for key,model,parent in [('areas',MaterialArea,None),('families',MaterialFamily,'area_id'),
        ('subfamilies',MaterialSubfamily,'family_id'),('types',MaterialProductType,'subfamily_id')]:
        result[key]=[{'id':row.id,'code':row.code,'name':row.name,**({parent:getattr(row,parent)} if parent else {})}
            for row in db.scalars(select(model).order_by(model.name)).all()]
    return result

@router.get('/policies')
def policies(db:Session=Depends(get_db)):
    return [policy_payload(row) for row in db.scalars(select(Policy).order_by(Policy.scope,Policy.scope_key)).all()]

@router.put('/policies')
def save_policy(body:PolicyInput,db:Session=Depends(get_db),user:User=Depends(operator)):
    validate_parameters(body.parameters)
    key=body.scope_key.strip()
    models={'AREA':MaterialArea,'FAMILY':MaterialFamily,'SUBFAMILY':MaterialSubfamily,'TYPE':MaterialProductType}
    if body.scope=='GLOBAL':key='*'
    elif body.scope in models and (not key.isdigit() or not db.get(models[body.scope],int(key))):
        raise HTTPException(422,'Clasificación no encontrada')
    elif body.scope=='DELEGATION' and key not in DELEGATIONS:raise HTTPException(422,'Delegación no admitida')
    elif body.scope=='ARTICLE_DELEGATION' and (len(key.rsplit('|',1))!=2 or key.rsplit('|',1)[1] not in DELEGATIONS or not key.rsplit('|',1)[0]):
        raise HTTPException(422,'Utiliza código de artículo|delegación, por ejemplo 03009|02')
    row=db.scalar(select(Policy).where(Policy.scope==body.scope,Policy.scope_key==key).with_for_update())
    previous=policy_payload(row) if row else None
    if row is None:
        row=Policy(scope=body.scope,scope_key=key,updated_by=user.id);db.add(row)
    row.parameters=body.parameters;row.active=body.active;row.updated_by=user.id
    db.flush();audit(db,user,'REPLENISHMENT_POLICY_SAVED',row.id,{'previous':previous,'new':policy_payload(row)})
    db.commit();return policy_payload(row)

def selected_product_codes(db,filters):
    query=select(Product.sku)
    for key,column in [('area_id',Product.material_area_id),('family_id',Product.material_family_id),
                       ('subfamily_id',Product.material_subfamily_id),('type_id',Product.material_product_type_id)]:
        if getattr(filters,key) is not None:query=query.where(column==getattr(filters,key))
    selected=list(db.scalars(query).all()) if any(getattr(filters,key) is not None for key in ('area_id','family_id','subfamily_id','type_id')) else None
    if filters.skus:selected=filters.skus if selected is None else sorted(set(selected).intersection(filters.skus))
    return selected

def classified_source(db,filters):
    selected=selected_product_codes(db,filters)
    limit=min(1000,module_settings(db)['article_limit'])
    source=read_replenishment_source(filters.date_from,filters.date_to,selected,limit)
    source['article_limit']=limit
    return enrich_source(db,source)

def enrich_source(db,source):
    codes=[row['sku'] for row in source['articles']]
    rows=[]
    for offset in range(0,len(codes),1000):
        rows.extend(db.execute(select(Product.sku,Product.material_area_id,Product.material_family_id,
            Product.material_subfamily_id,Product.material_product_type_id).where(Product.sku.in_(codes[offset:offset+1000]))).all())
    metadata={row[0]:dict(zip(('area_id','family_id','subfamily_id','type_id'),row[1:])) for row in rows}
    names={key:{row.id:row.name for row in db.scalars(select(model)).all()} for key,model in
        [('area',MaterialArea),('family',MaterialFamily),('subfamily',MaterialSubfamily),('type',MaterialProductType)]}
    for article in source['articles']:
        article.update(metadata.get(article['sku'],{}))
        for key in names:article[key]=names[key].get(article.get(key+'_id'))
    return source

def job_payload(db,row):
    result=scenario_payload(db.get(Scenario,row.scenario_id)) if row.status=='COMPLETED' else None
    progress=100 if row.status=='COMPLETED' else min(99,int(row.processed_articles*100/max(1,row.total_articles)))
    return {'id':row.public_id,'status':row.status,'phase':row.phase,'progress':progress,
        'total_articles':row.total_articles,'processed_articles':row.processed_articles,
        'total_batches':row.total_batches,'completed_batches':row.completed_batches,'batch_size':1000,
        'started_at':row.started_at,'finished_at':row.finished_at,'error':row.error,'result':result,
        'filters':row.request['filters'],'temporary_parameters':row.request['temporary_parameters']}

@router.post('/jobs',status_code=202)
def enqueue(body:RunInput,db:Session=Depends(get_db),user:User=Depends(operator)):
    validate_parameters(body.temporary_parameters)
    if body.source_id or body.preview_only:
        raise HTTPException(422,'El cálculo por bloques consulta y guarda sus propios datos de EXIT.')
    # Comparables require the complete population; the current simple calculation never uses them.
    if body.temporary_parameters.get('simple_formula') is not True:
        raise HTTPException(422,'El cálculo por bloques utiliza la fórmula simple visible de reposición.')
    db.scalar(select(User).where(User.id==user.id).with_for_update())
    existing=db.scalar(select(Job).where(Job.created_by==user.id,Job.status.in_(['QUEUED','RUNNING'])).order_by(Job.id))
    if existing:return job_payload(db,existing)
    row=Job(created_by=user.id,request=body.model_dump(mode='json'),article_limit=module_settings(db)['article_limit'])
    db.add(row);db.flush()
    audit(db,user,'REPLENISHMENT_QUEUED',row.public_id,{'batch_size':1000,'article_limit':row.article_limit})
    db.commit();return job_payload(db,row)

@router.get('/jobs/current')
def current_job(db:Session=Depends(get_db),user:User=Depends(operator)):
    latest=db.scalar(select(Job).where(Job.created_by==user.id).order_by(Job.id.desc()).limit(1))
    completed=db.scalar(select(Job).where(Job.created_by==user.id,Job.status=='COMPLETED').order_by(Job.id.desc()).limit(1))
    return {'job':job_payload(db,latest) if latest else None,
        'latest_result':scenario_payload(db.get(Scenario,completed.scenario_id)) if completed else None}

@router.get('/jobs/{job_id}')
def read_job(job_id:str,db:Session=Depends(get_db),user:User=Depends(operator)):
    row=db.scalar(select(Job).where(Job.public_id==job_id,Job.created_by==user.id))
    if not row:raise HTTPException(404,'Cálculo no encontrado')
    return job_payload(db,row)

@router.post('/simulations',status_code=201)
def run_simulation(body:RunInput,db:Session=Depends(get_db),user:User=Depends(operator)):
    validate_parameters(body.temporary_parameters)
    filters=body.filters.model_dump(mode='json')
    if body.source_id:
        base=scenario_or_404(db,body.source_id)
        if base.filters!=filters:raise HTTPException(422,'Para reutilizar datos conserva los filtros originales o actualiza la extracción')
        source=base.source_snapshot
        if source.get('batched'):raise HTTPException(422,'Para repetir un cálculo por bloques utiliza /jobs.')
    else:
        try:source=classified_source(db,body.filters)
        except ValueError as exc:raise HTTPException(422,str(exc)) from exc
        except Exception as exc:
            logger.exception('Error leyendo fuentes de reposición EXIT')
            raise HTTPException(503,'No se pudieron leer las fuentes EXIT indicadas. Revisa el registro de la API y la conexión') from exc
    if len(source['articles'])>1000:
        raise HTTPException(422,'Para más de 1.000 artículos utiliza el cálculo por bloques (/jobs).')
    snapshots=[]  # El cálculo simple usa exclusivamente los parámetros visibles enviados.
    fingerprint=hashlib.sha256(json.dumps(source,sort_keys=True,ensure_ascii=False).encode()).hexdigest()
    scenario=Scenario(created_by=user.id,source_snapshot=source,source_fingerprint=fingerprint,filters=filters,
        policies_snapshot=snapshots,temporary_parameters=body.temporary_parameters,summary={})
    db.add(scenario);db.flush()
    results=[]
    if not body.preview_only:
        try:
            context=prepare_context(source,body.filters.delegations,snapshots,body.temporary_parameters)
            results=[calculate_article(article,delegation,context) for article in source['articles'] for delegation in body.filters.delegations]
        except ValueError as exc:raise HTTPException(422,f'Parámetros incompatibles al resolver políticas: {exc}') from exc
        annotate_central(results,source);scenario.summary=summarize(results)
        db.add_all([Result(scenario_id=scenario.id,sku=row['sku'],delegation=row['delegation'],payload=row) for row in results])
    audit(db,user,'REPLENISHMENT_SIMULATED',scenario.public_id,{'preview_only':body.preview_only,'source_fingerprint':fingerprint})
    db.commit();return scenario_payload(scenario)

class ResultFilters(BaseModel):
    q:str=Field(default='',max_length=200)
    alerts_only:bool=False
    needed_only:bool=False
    central_only:bool=False
    overstock_only:bool=False
    consumption:Literal['ALL','WITH','WITHOUT']='ALL'
    minimum_monthly_consumption:float=Field(0,ge=0,le=1e9,allow_inf_nan=False)
    sort:Literal['sku','quantity','coverage','alerts']='sku'
    descending:bool=False


def filtered_results(scenario_id,filters):
    q=filters.q;alerts_only=filters.alerts_only;needed_only=filters.needed_only
    central_only=filters.central_only;overstock_only=filters.overstock_only;consumption=filters.consumption
    query=select(Result).where(Result.scenario_id==scenario_id)
    if q:
        query=query.where(or_(func.lower(Result.sku).contains(q.lower(),autoescape=True),
            func.lower(Result.payload['description'].as_string()).contains(q.lower(),autoescape=True)))
    if alerts_only:query=query.where(func.json_array_length(Result.payload['recommendations'])>0)
    if needed_only:query=query.where(Result.payload['final_quantity'].as_float()>0)
    if central_only:query=query.where(Result.payload['central_can_supply'].as_boolean().is_(True))
    if overstock_only:query=query.where(Result.payload['overstock'].as_boolean().is_(True))
    if filters.minimum_monthly_consumption>0:
        monthly=func.coalesce(Result.payload['monthly_mean'].as_float(),Result.payload['daily_mean'].as_float()*30)
        query=query.where(monthly>=filters.minimum_monthly_consumption)
    if consumption=='WITHOUT':query=query.where(Result.payload['total_demand'].as_float()<=0)
    if consumption=='WITH':query=query.where(Result.payload['total_demand'].as_float()>0)
    return query


def ordered_results(query,sort,descending):
    order={'sku':Result.sku,'quantity':Result.payload['final_quantity'].as_float(),
        'coverage':Result.payload['coverage_days'].as_float(),'alerts':func.json_array_length(Result.payload['recommendations'])}[sort]
    return query.order_by(order.desc() if descending else order.asc(),Result.sku,Result.delegation)


def result_summary(db,query,full=False):
    data=query.subquery()
    payload=data.c.payload
    count_if=lambda condition:func.coalesce(func.sum(case((condition,1),else_=0)),0)
    columns=[func.count().label('article_delegations'),func.count(func.distinct(data.c.sku)).label('articles_analyzed'),
        count_if(payload['final_quantity'].as_float()>0).label('to_replenish'),
        count_if(payload['central_can_supply'].as_boolean().is_(True)).label('central_covered'),
        count_if(payload['overstock'].as_boolean().is_(True)).label('overstock'),
        count_if(payload['total_demand'].as_float()<=0).label('without_movement')]
    if full:columns.extend([
        func.coalesce(func.sum(payload['final_quantity'].as_float()),0).label('units_proposed'),
        count_if(payload['irregular'].as_boolean().is_(True)).label('irregular'),
        func.coalesce(func.sum(func.json_array_length(payload['recommendations'])),0).label('alerts'),
        func.coalesce(func.sum(payload['target_stock'].as_float()),0).label('target_stock'),
        func.coalesce(func.avg(case((payload['daily_mean'].as_float()>0,payload['target_stock'].as_float()/func.nullif(payload['daily_mean'].as_float(),0)),else_=None)),0).label('average_coverage_days')])
    row=db.execute(select(*columns)).mappings().one()
    return dict(row)


def result_scenario(db,public_id):
    row=db.execute(select(Scenario.id,Scenario.public_id,Scenario.version,Scenario.filters,Scenario.temporary_parameters,
        Scenario.policies_snapshot,Scenario.source_fingerprint,Scenario.source_snapshot['captured_at'].as_string().label('captured_at'))
        .where(Scenario.public_id==public_id)).one_or_none()
    if row is None:raise HTTPException(404,'Cálculo no encontrado')
    ensure_published(db,row.id)
    return row


@router.get('/simulations/{scenario_id}/results')
def results(scenario_id:str,filters:ResultFilters=Depends(),page:int=Query(1,ge=1),
        page_size:int=Query(50,ge=1,le=100),db:Session=Depends(get_db)):
    scenario=result_scenario(db,scenario_id)
    query=filtered_results(scenario.id,filters)
    summary=result_summary(db,query)
    rows=db.scalars(ordered_results(query,filters.sort,filters.descending).offset((page-1)*page_size).limit(page_size)).all()
    return {'items':[{'id':row.id,**row.payload} for row in rows],'total':summary['article_delegations'],
        'summary':summary,'page':page,'version':scenario.version}


def source_for_result(db,scenario,result):
    if not scenario.source_snapshot.get('batched'):return scenario.source_snapshot
    source=db.scalar(select(Batch.source).join(Job,Job.id==Batch.job_id).where(
        Job.scenario_id==scenario.id,Batch.number==result.payload.get('source_batch')))
    if not source:raise HTTPException(404,'Bloque de origen no encontrado')
    # The simple formula is independent per SKU; detail/adjustments need no other article histories.
    return {**source,**{key:[row for row in source[key] if row['sku']==result.sku] for key in ('articles','sales','stocks')}}

def article_context(scenario,result,db):
    source=source_for_result(db,scenario,result)
    article=next((row for row in source['articles'] if row['sku']==result.sku),None)
    if not article:raise HTTPException(404,'Artículo ausente en la instantánea')
    context=prepare_context(source,scenario.filters['delegations'],scenario.policies_snapshot,scenario.temporary_parameters)
    return article,context

@router.get('/simulations/{scenario_id}/results/{result_id}')
def detail(scenario_id:str,result_id:int,db:Session=Depends(get_db)):
    scenario=scenario_or_404(db,scenario_id);row=result_or_404(db,scenario,result_id)
    article,context=article_context(scenario,row,db)
    calculated=calculate_article(article,row.delegation,context,row.payload.get('manual_parameters'),True)
    history=db.scalars(select(Adjustment).where(Adjustment.result_id==row.id).order_by(Adjustment.created_at)).all()
    return {'id':row.id,**row.payload,'daily_sales':calculated['daily_sales'],'scenario_version':scenario.version,
        'adjustments':[{'created_at':event.created_at,'user_id':event.user_id,'reason':event.reason,'changes':event.changes} for event in history]}

@router.post('/simulations/{scenario_id}/results/{result_id}/recalculate')
def recalculate(scenario_id:str,result_id:int,body:RecalculateInput,db:Session=Depends(get_db)):
    validate_parameters(body.parameters)
    scenario=scenario_or_404(db,scenario_id);row=result_or_404(db,scenario,result_id)
    article,context=article_context(scenario,row,db)
    try:result=calculate_article(article,row.delegation,context,{**row.payload.get('manual_parameters',{}),**body.parameters},True)
    except ValueError as exc:raise HTTPException(422,str(exc)) from exc
    other=[dict(item.payload) for item in db.scalars(select(Result).where(Result.scenario_id==scenario.id,Result.sku==row.sku,Result.id!=row.id)).all()]
    annotate_central([*other,result],context['source'])
    return result

@router.post('/simulations/{scenario_id}/results/{result_id}/adjustments')
def adjust(scenario_id:str,result_id:int,body:AdjustmentInput,db:Session=Depends(get_db),user:User=Depends(operator)):
    validate_parameters(body.parameters)
    scenario=scenario_or_404(db,scenario_id,True);row=result_or_404(db,scenario,result_id)
    if scenario.version!=body.expected_version:raise HTTPException(409,'Otra persona modificó el escenario. Actualiza los datos')
    article,context=article_context(scenario,row,db);previous=dict(row.payload)
    patch={**previous.get('manual_parameters',{}),**body.parameters}
    try:updated=calculate_article(article,row.delegation,context,patch)
    except ValueError as exc:raise HTTPException(422,str(exc)) from exc
    if 'source_batch' in previous:updated['source_batch']=previous['source_batch']
    updated['manual_parameters']=patch
    if not body.excluded and body.quantity is not None and body.quantity>0 and updated['parameters']['multiple_mode']=='EXIT':
        multiple=updated['effective_multiple']
        if multiple is None:raise HTTPException(422,'No hay un EX_multiplo válido. Corrige EXIT o selecciona otro modo de agrupación')
        if abs(rounded_dispatch(body.quantity,multiple)-body.quantity)>1e-8:
            raise HTTPException(422,f'La cantidad ajustada debe ser múltiplo de {multiple:g}, según EX_multiplo de EXIT')
    updated['adjusted_quantity']=0 if body.excluded else body.quantity
    updated['final_quantity']=updated['recommended_quantity'] if updated['adjusted_quantity'] is None else updated['adjusted_quantity']
    updated['manually_excluded']=body.excluded;updated['calculated_quantity']=previous['calculated_quantity']
    row.payload=updated
    all_rows=db.scalars(select(Result).where(Result.scenario_id==scenario.id,Result.sku==row.sku)).all()
    payloads=[dict(item.payload) for item in all_rows];annotate_central(payloads,context['source'])
    for item,payload in zip(all_rows,payloads):item.payload=payload
    db.flush()
    scenario.summary=result_summary(db,filtered_results(scenario.id,ResultFilters()),full=True);scenario.version+=1
    db.add(Adjustment(scenario_id=scenario.id,result_id=row.id,user_id=user.id,reason=body.reason,previous=previous,changes=body.model_dump()))
    audit(db,user,'REPLENISHMENT_MANUAL_ADJUSTMENT',scenario.public_id,{'sku':row.sku,'delegation':row.delegation,'reason':body.reason})
    db.commit();return {'scenario':scenario_payload(scenario),'result':{'id':row.id,**row.payload}}

@router.get('/simulations/{scenario_id}/export.xlsx')
def export(scenario_id:str,filters:ResultFilters=Depends(),db:Session=Depends(get_db),mode:Literal["simple","expanded"]="simple"):
    scenario=result_scenario(db,scenario_id)
    query=ordered_results(filtered_results(scenario.id,filters),filters.sort,filters.descending)
    rows=db.scalars(query.execution_options(yield_per=500))
    book=Workbook(write_only=True);sheet=book.create_sheet('Reposición')
    sheet.freeze_panes='A2'
    def append(sheet,values):
        cells=[]
        for value in values:
            if isinstance(value,str):
                value=ILLEGAL_CHARACTERS_RE.sub('',value)
                cell=WriteOnlyCell(sheet,value=value);cell.data_type='s';cells.append(cell)
            else:cells.append(value)
        sheet.append(cells)
    fields=[('sku','Código'),('description','Descripción'),('delegation','Delegación'),('unit_type','Tipo unidad'),('family','Familia'),('target_mode','Objetivo elegido'),
        ('central_stock','Stock Almeiras'),('central_outgoing','Comprometido Almeiras'),('central_incoming','Pendiente recibir Almeiras'),('central_monthly_consumption','Consumo mensual Almeiras'),('central_reserve','Reserva Almeiras'),('central_available','Disponible para despacho Almeiras'),('central_total_requested','Propuesta conjunta'),('central_can_supply','Almeiras cubre propuesta conjunta'),('overstock','Sobrestock'),('excess_stock','Exceso previsto'),('no_consumption','Sin consumo'),('stock','Stock'),('incoming','Pendiente recibir'),('outgoing','Pendiente servir'),('inventory_position','Disponible previsto'),('monthly_mean','Consumo mensual medio'),
        ('safety_stock','Seguridad'),('minimum_stock','Mínimo'),('maximum_stock','Stock máximo'),('expected_cycle_consumption','Consumo hasta próxima reposición'),('replenishment_point','Punto de reposición'),('target_stock','Objetivo aplicado'),('calculated_quantity','Cantidad calculada'),
        ('exit_multiple','Múltiplo EXIT'),('effective_multiple','Múltiplo aplicado'),('quantity_before_rounding','Cantidad antes de agrupar'),('recommended_quantity','Cantidad recomendada'),('adjusted_quantity','Ajuste manual'),('final_quantity','Cantidad final'),('method','Método'),('confidence','Confianza')]
    if mode=='simple':
        fields=[('sku','Código'),('description','Descripción'),('delegation','Delegación'),
            ('stock','Stock'),('incoming','Por recibir'),('outgoing','Por servir'),
            ('inventory_position','Disponible previsto'),('unit_type','Unidad'),('monthly_mean','Consumo mensual'),
            ('minimum_stock','Mínimo'),('replenishment_point','Punto reposición'),
            ('maximum_stock','Stock máximo'),('final_quantity','A reponer'),
            ('central_available','Disponible para despacho'),('central_status','Cobertura Almeiras'),
            ('notices','Avisos')]
    extra=['Reglas aplicadas','Recomendaciones'] if mode=='expanded' else []
    append(sheet,[label for _,label in fields]+extra)
    count=0
    for result in rows:
        row={**result.payload,'monthly_mean':result.payload.get('monthly_mean',result.payload['daily_mean']*30)}
        row['central_status']=('Sin necesidad' if not row['final_quantity'] else 'Revisar datos' if row.get('central_available') is None else 'Cubre conjunto' if row.get('central_can_supply') else 'No cubre conjunto')
        row['notices']='; '.join(item['reason'] for item in row['recommendations'])
        if mode=='simple':
            for key in ('minimum_stock','replenishment_point','maximum_stock'):
                if row.get(key) is not None:row[key]=math.ceil(row[key])
            row['monthly_mean']=round(row['monthly_mean'],2)
            if row.get('central_available') is not None:row['central_available']=math.floor(row['central_available'])
        values=[row.get(key) for key,_ in fields]
        if mode=='expanded':values+=[' → '.join(row['applied_rules']),row['notices']]
        append(sheet,values);count+=1
    sheet.auto_filter.ref=f'A1:{get_column_letter(len(fields)+len(extra))}{count+1}'
    metadata=book.create_sheet('Contexto') if mode=='expanded' else None
    if metadata is not None:append(metadata,['Dato','Valor'])
    for key,value in [('Cálculo',scenario.public_id),('Corte',scenario.captured_at),('Huella de datos',scenario.source_fingerprint),
        ('Filtros del resultado',json.dumps(filters.model_dump(),ensure_ascii=False)),('Filas exportadas',count),('Filtros del cálculo',json.dumps(scenario.filters,ensure_ascii=False)),('Parámetros temporales',json.dumps(scenario.temporary_parameters,ensure_ascii=False)),
        ('Políticas',json.dumps(scenario.policies_snapshot,ensure_ascii=False)),('Modo','Solo cálculo; no genera pedidos. Almeiras protegido por su consumo y reserva.')]:
        if metadata is not None:append(metadata,[key,value])
    buffer=io.BytesIO();book.save(buffer)
    return Response(buffer.getvalue(),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':'attachment; filename="reposicion.xlsx"'})
