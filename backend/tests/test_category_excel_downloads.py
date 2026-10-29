import io
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select
from app.bootstrap import main as bootstrap
from app.main import app
from app.db import SessionLocal
from app.models import Product, ProductImageMissing
from app.category_import import SOURCE_SHEET, TEMPLATE_HEADERS, parse_category_file

bootstrap()
PREFIX='/api/v1/store/catalog-categories'
MEDIA='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def login(email='operador@bermudez.test'):
    client=TestClient(app)
    assert client.post('/api/v1/auth/login',json={'email':email,'password':'123456'}).status_code==200
    return client


def workbook(response,filename):
    assert response.status_code==200,response.text[:200] if response.status_code!=200 else ''
    assert response.headers['content-type']==MEDIA
    assert filename in response.headers['content-disposition']
    assert response.content[:2]==b'PK'
    book=load_workbook(io.BytesIO(response.content))
    assert book.active.freeze_panes=='A2'
    assert book.active.auto_filter.ref
    return book


def test_full_classification_is_real_excel_preserves_codes_and_literal_descriptions(monkeypatch):
    monkeypatch.setattr('app.category_import.fetch_all_catalog_articles_for_export',lambda:[
        {'code':'00123','description':'=SUM(1,2)\x0b','inactive':'1','manufacturer_reference':'+REF-001'},
        {'code':'00007','description':'VÁLVULA; LATÓN','inactive':'0','unit':'ROL'},
    ])
    sheet=workbook(login().get(PREFIX+'/export.xlsx'),'clasificacion_articulos.xlsx').active
    headers=[cell.value for cell in sheet[1]]
    rows={row[0].value:dict(zip(headers,[cell.value for cell in row])) for row in list(sheet)[1:]}
    assert len(headers)==25 and sheet.max_row==3
    assert set(rows)=={'00123','00007'}
    assert rows['00123']['Descripcion']=='=SUM(1,2)' and sheet['B2'].data_type=='s'
    assert rows['00123']['ReferenciaFabricanteEXIT']=='+REF-001'
    assert rows['00123']['Activo']=='No' and rows['00123']['EstadoClasificacion']=='SIN_CLASIFICAR'
    assert rows['00007']['Activo']=='Sí' and rows['00007']['UnidadMedidaEXIT']=='ROL'
    assert sheet['A2'].data_type=='s' and sheet['A2'].number_format=='@'


def test_full_excel_preserves_complete_exit_scope_and_local_classification(monkeypatch):
    with SessionLocal() as db:
        product=db.scalar(select(Product).order_by(Product.id));code=product.sku
    monkeypatch.setattr('app.category_import.fetch_all_catalog_articles_for_export',lambda:[
        {'code':code,'description':'Maestra EXIT','inactive':'1'}, {'code':'000-EXIT-ONLY','description':'Solo EXIT','inactive':'1'}])
    client=login();sheet=workbook(client.get(PREFIX+'/export.xlsx'),'clasificacion_articulos.xlsx').active
    headers=[cell.value for cell in sheet[1]]
    rows={row[0].value:dict(zip(headers,[cell.value for cell in row])) for row in list(sheet)[1:]}
    assert set(rows)=={code,'000-EXIT-ONLY'}
    assert rows[code]['ExisteEnWeb']=='Sí' and rows['000-EXIT-ONLY']['ExisteEnWeb']=='No'
    assert rows[code]['Descripcion']=='Maestra EXIT'


def test_excel_template_is_editable_and_accepted_by_existing_import_parser():
    response=login().get(PREFIX+'/template.xlsx')
    sheet=workbook(response,'plantilla_carga_categorias.xlsx').active
    assert sheet.title==SOURCE_SHEET
    assert [cell.value for cell in sheet[1]]==TEMPLATE_HEADERS
    assert sheet['A2'].value=='01234' and sheet['A2'].data_type=='s'
    rows,title=parse_category_file('plantilla.xlsx',response.content)
    assert title==SOURCE_SHEET and rows[0]['CodigoArticulo']=='01234'


def test_unclassified_excel_matches_import_columns_and_preserves_reference_text(monkeypatch):
    with SessionLocal() as db:
        product=db.scalar(select(Product).order_by(Product.id));pk=product.id
        old=(product.source_system,product.classification_status)
        product.source_system='EXIT_CATALOG';product.classification_status='SIN_CLASIFICAR';code=product.sku;db.commit()
    monkeypatch.setattr('app.category_import.fetch_catalog_articles_by_codes',lambda *a,**kw:{code:{'brand_name':'=MARCA','manufacturer_reference':'000-REF','unit':'ROL'}})
    try:
        response=login().get(PREFIX+'/missing.xlsx')
        sheet=workbook(response,'articulos_sin_clasificacion.xlsx').active
        headers=[cell.value for cell in sheet[1]]
        assert headers==TEMPLATE_HEADERS+['EstadoVinculo','Motivo']
        row=next(row for row in sheet.iter_rows(min_row=2) if row[0].value==code)
        values=dict(zip(headers,[cell.value for cell in row]))
        assert values['Marca']=='=MARCA' and row[5].data_type=='s'
        assert values['ReferenciaFabricante']=='000-REF' and values['UnidadMedida']=='ROL'
        parsed,title=parse_category_file('sin_clasificar.xlsx',response.content)
        assert title==SOURCE_SHEET and any(row['CodigoArticulo']==code for row in parsed)
    finally:
        with SessionLocal() as db:
            product=db.get(Product,pk);product.source_system,product.classification_status=old;db.commit()


def test_photographs_download_is_excel_with_literal_names():
    with SessionLocal() as db:
        product=db.scalar(select(Product).order_by(Product.id));pk=product.id;old_name=product.short_description
        product.short_description='=NOMBRE\x0b'
        missing=db.scalar(select(ProductImageMissing).where(ProductImageMissing.product_id==pk))
        created=missing is None
        if missing is None:
            missing=ProductImageMissing(product_id=pk,sku=product.sku);db.add(missing)
        old=(missing.status,missing.reason);missing.status='PENDIENTE';missing.reason='@revisar';code=product.sku;db.commit()
    try:
        sheet=workbook(login().get('/api/v1/store/catalog-images/missing.xlsx'),'imagenes_pendientes.xlsx').active
        headers=[cell.value for cell in sheet[1]]
        assert headers==['codigo_material','nombre_actual','motivo','estado','image_url','image_path','proveedor','aprobado']
        row=next(row for row in sheet.iter_rows(min_row=2) if row[0].value==code)
        assert row[1].value=='=NOMBRE' and row[1].data_type=='s'
        assert row[2].value=='@revisar' and row[7].value=='NO'
    finally:
        with SessionLocal() as db:
            db.get(Product,pk).short_description=old_name
            row=db.scalar(select(ProductImageMissing).where(ProductImageMissing.product_id==pk))
            if created:db.delete(row)
            else:row.status,row.reason=old
            db.commit()


@pytest.mark.parametrize('path',[PREFIX+'/export.xlsx',PREFIX+'/missing.xlsx',PREFIX+'/template.xlsx','/api/v1/store/catalog-images/missing.xlsx'])
def test_excel_downloads_require_operator_permissions(path):
    assert TestClient(app).get(path).status_code==401
    assert login('compras001@cliente.test').get(path).status_code==403


def test_export_excel_empty_master_and_source_failure_do_not_return_partial_data(monkeypatch):
    client=login()
    monkeypatch.setattr('app.category_import.fetch_all_catalog_articles_for_export',lambda:[])
    assert workbook(client.get(PREFIX+'/export.xlsx'),'clasificacion_articulos.xlsx').active.max_row==1
    def unavailable():raise RuntimeError('EXIT unavailable')
    monkeypatch.setattr('app.category_import.fetch_all_catalog_articles_for_export',unavailable)
    response=client.get(PREFIX+'/export.xlsx')
    assert response.status_code==503 and 'parcial' in response.json()['detail']
