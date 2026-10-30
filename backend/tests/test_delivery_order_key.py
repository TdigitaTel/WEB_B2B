from datetime import datetime,timezone
from contextlib import contextmanager
from types import SimpleNamespace
import pytest
from fastapi.testclient import TestClient
from app import main as module,erp_db
from app.bootstrap import main as bootstrap
from app.db import SessionLocal
from app.models import Order
from app.delivery_links import matches_note

bootstrap()
pytestmark=pytest.mark.usefixtures('operator_with_almeiras_access')


def test_requires_all_four_order_components():
    order=SimpleNamespace(company_code=1,exit_company_code=1,exit_order_year=2026,exit_series='AL',exit_order_number='8811',nro_pedido_exit='2026/AL/8811')
    note={'company_code':1,'order_number':'2026/AL/8811'}
    assert matches_note(note,order)
    assert not matches_note({'order_number':'2026/AL/8811'},order)
    for reference in ['2025/AL/8811','2026/OT/8811','2026/AL/8812','8811']:
        assert not matches_note({**note,'order_number':reference},order)
    assert not matches_note({**note,'company_code':2},order)
    order.exit_order_year=None;order.exit_series=None;order.exit_order_number=None;order.nro_pedido_exit='8811'
    assert not matches_note(note,order)


def test_batch_sql_filters_company_and_processes_more_than_one_thousand(monkeypatch):
    calls=[]
    class Cursor:
        def __enter__(self):return self
        def __exit__(self,*args):pass
        def execute(self,sql,params):calls.append((sql,params))
        def fetchall(self):return []
    class Connection(Cursor):
        def cursor(self):return Cursor()
    monkeypatch.setattr(erp_db,'connect_sqlserver',lambda:Connection())
    monkeypatch.setattr(erp_db,'_delivery_registration_sql',lambda connection,sql:sql)
    keys=[(i%2+1,2026,'AL',str(i)) for i in range(1001)]
    assert erp_db.fetch_delivery_statuses_for_orders(keys)==[]
    assert [len(params) for _,params in calls]==[1600,1600,804]
    assert all('CodigoEmpresa=%s AND EjercicioPedido=%s' in sql for sql,_ in calls)
    assert tuple(calls[0][1][:4])==keys[0]
    assert len({tuple(params[i:i+4]) for _,params in calls for i in range(0,len(params),4)})==1001


def test_delivery_document_lines_do_not_mix_companies(monkeypatch):
    stamp=datetime.now(timezone.utc)
    headers=[dict(CodigoEmpresa=company,EjercicioAlbaran=2026,SerieAlbaran='AL',NumeroAlbaran=90,FechaAlbaran=stamp,
        EjercicioPedido=2026,SeriePedido='AL',NumeroPedido=8811) for company in [1,2]]
    lines=[dict(CodigoEmpresa=company,EjercicioAlbaran=2026,SerieAlbaran='AL',NumeroAlbaran=90,CodigoArticulo=f'COMPANY-{company}') for company in [1,2]]
    class Cursor:
        def __enter__(self):return self
        def __exit__(self,*args):pass
        def execute(self,sql,params):self.sql=sql
        def fetchall(self):return lines if 'WITH selected' in self.sql else headers
    class Connection(Cursor):
        def cursor(self):return Cursor()
    monkeypatch.setattr(erp_db,'connect_sqlserver',lambda:Connection())
    documents=erp_db.fetch_customer_delivery_notes('CLIENT')
    assert [doc['items'][0]['sku'] for doc in documents]==['COMPANY-1','COMPANY-2']
    assert [doc['exit_company_code'] for doc in documents]==[1,2]
    assert all(doc['exit_order_year']==2026 and doc['exit_series']=='AL' and doc['exit_order_number']=='8811' for doc in documents)


def test_operator_cannot_open_note_with_wrong_or_missing_company(monkeypatch):
    with SessionLocal() as db:
        order=Order(order_number='WEB-NOTE-4-KEY',company_code=0,delegation_code='00',customer_code='CLI00001',origen_pedido='B2B',estado_registro_exit='REGISTRADO',nro_pedido_exit='2026/AL/8811',exit_order_year=2026,exit_series='AL',exit_order_number='8811')
        db.add(order);db.commit();oid=order.public_id
    document={'number':'2026/AL/9900','company_code':1,'order_number':'2026/AL/8811','items':[]}
    monkeypatch.setattr(module,'fetch_customer_delivery_notes',lambda *args:[document])
    operator=TestClient(module.app);operator.post('/api/v1/auth/login',json={'email':'operador@bermudez.test','password':'123456'})
    path=f'/api/v1/store/orders/{oid}/delivery-note?number=2026/AL/9900'
    assert operator.get(path).status_code==404
    document['company_code']=0
    assert operator.get(path).status_code==200
    del document['company_code']
    assert operator.get(path).status_code==404


def test_foreign_company_note_does_not_change_customer_order_state(monkeypatch):
    with SessionLocal() as db:
        order=Order(order_number='WEB-NOTE-4-STATE',company_code=0,delegation_code='00',customer_code='CLI00001',origen_pedido='B2B',estado_registro_exit='REGISTRADO',nro_pedido_exit='2026/AL/8812',exit_order_year=2026,exit_series='AL',exit_order_number='8812')
        db.add(order);db.commit();oid=order.public_id
    note={'company_code':1,'order_number':'2026/AL/8812','delivery_number':'2026/AL/9901','attended_at':datetime.now(timezone.utc),'is_printed':False,'is_invoiced':False}
    monkeypatch.setattr(module,'fetch_customer_exit_orders',lambda *args:[])
    monkeypatch.setattr(module,'fetch_customer_delivery_statuses',lambda *args:[note])
    customer=TestClient(module.app);customer.post('/api/v1/auth/login',json={'email':'compras001@cliente.test','password':'123456'})
    foreign=next(row for row in customer.get('/api/v1/orders?state=TODOS').json() if row['id']==oid)
    assert foreign['estado_registro_exit']=='REGISTRADO' and not foreign['delivery_notes']
    note['company_code']=0
    correct=next(row for row in customer.get('/api/v1/orders?state=TODOS').json() if row['id']==oid)
    assert correct['estado_registro_exit']=='ATENDIDO' and len(correct['delivery_notes'])==1


def test_customer_open_and_export_keep_the_order_context(monkeypatch):
    from io import BytesIO
    from openpyxl import load_workbook
    with SessionLocal() as db:
        order=Order(order_number='WEB-NOTE-CONTEXT',company_code=0,delegation_code='00',customer_code='CLI00001',origen_pedido='B2B',estado_registro_exit='REGISTRADO',exit_order_year=2026,exit_series='AL',exit_order_number='8820')
        foreign=Order(order_number='WEB-NOTE-FOREIGN',company_code=0,delegation_code='00',customer_code='OTHER-CUSTOMER',origen_pedido='B2B',estado_registro_exit='REGISTRADO',exit_order_year=2026,exit_series='AL',exit_order_number='8821')
        db.add_all([order,foreign]);db.commit();oid,foreign_id=order.public_id,foreign.public_id
    rows=[{'id':f'{company}:note','number':'2026-AL-9920','company_code':company,'order_number':'2026/AL/8820','created_at':datetime(2026,10,29),'items':[],'total':1,'subtotal':1,'tax_total':0} for company in [0,1]]
    monkeypatch.setattr(module,'fetch_customer_delivery_notes',lambda *args:rows)
    customer=TestClient(module.app);customer.post('/api/v1/auth/login',json={'email':'compras001@cliente.test','password':'123456'})
    response=customer.get('/api/v1/delivery-notes',params={'order_id':oid,'document_number':'2026-AL-9920'})
    assert response.status_code==200,response.text
    assert [row['company_code'] for row in response.json()]==[0]
    assert customer.get('/api/v1/delivery-notes',params={'order_id':foreign_id}).status_code==404
    excel=customer.get('/api/v1/documents/export',params={'kind':'ALBARAN','order_id':oid})
    assert excel.status_code==200,excel.text if excel.status_code!=200 else ''
    assert load_workbook(BytesIO(excel.content)).active.max_row==2
